#!/usr/bin/env python3
"""Книга по доходным домам: три лота, колонки под этот тип объекта.

    python3 -m reportgen.fib_book --rows fib/rows.json --out corner_houses.xlsx

Колонки не те, что у квартир. Здесь покупают не жильё, а поток: важно, что
внизу, что сверху, сдано ли, и почём. Поэтому «Аренда, €/мес» и «€/м² аренды»
уступают место колонкам «Первый этаж», «Выше», «Доход, €/год» и «Доходность».

Половина этих чисел лежит только в прозе объявления, а не в полях. Поэтому у
каждой такой ячейки есть пометка в примечании: откуда взято. Пустая ячейка
означает «в объявлении не сказано» — и это честнее, чем правдоподобная цифра.
"""
import argparse, json, os, sys

COLS = [
    ("№", 5), ("Город", 11), ("Адрес", 30), ("Что это", 30),
    ("Угловой", 22), ("Первый этаж", 40), ("Выше", 34),
    ("Год", 11), ("Площадь, м²", 12), ("Цена, €", 12), ("€/м²", 9),
    ("Доход, €/год", 13), ("Доходность", 11), ("Состояние", 30), ("Ссылка", 16),
]


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--rows", default="fib/rows.json")
    ap.add_argument("--out", default="corner_houses.xlsx")
    a = ap.parse_args()

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    rows = json.load(open(a.rows, encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Дома"

    ink = "1F2933"
    head_fill = PatternFill("solid", fgColor="1F2933")
    band = PatternFill("solid", fgColor="F4F6F8")
    thin = Side(style="thin", color="D6DCE4")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    for i, (title, w) in enumerate(COLS, 1):
        c = ws.cell(1, i, title)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 34

    for n, r in enumerate(rows, 1):
        vals = [n, r.get("city"), r.get("address"), r.get("what"), r.get("corner"),
                r.get("ground"), r.get("above"), r.get("year"), r.get("area"),
                r.get("price"), None, r.get("income"), r.get("yield"),
                r.get("condition"), "funda in business"]
        for i, v in enumerate(vals, 1):
            c = ws.cell(n + 1, i, v)
            c.border = box
            c.alignment = Alignment(vertical="top", wrap_text=i in (3, 4, 5, 6, 7, 14))
            if n % 2 == 0:
                c.fill = band
        # €/м² считает книга, а не я: формула видна и пересчитается при правке
        area_c, price_c = f"I{n + 1}", f"J{n + 1}"
        ws.cell(n + 1, 11).value = f'=IF(AND(ISNUMBER({area_c}),ISNUMBER({price_c})),ROUND({price_c}/{area_c},0),"")'
        ws.cell(n + 1, 10).number_format = '#,##0 "€"'
        ws.cell(n + 1, 11).number_format = '#,##0'
        ws.cell(n + 1, 12).number_format = '#,##0 "€"'
        ws.cell(n + 1, 9).number_format = '#,##0'
        link = r.get("url")
        if link:
            c = ws.cell(n + 1, 15)
            c.hyperlink = link
            c.font = Font(color="1155CC", underline="single")
        for key, col in (("note_area", 9), ("note_income", 12), ("note_year", 8)):
            if r.get(key):
                from openpyxl.comments import Comment
                ws.cell(n + 1, col).comment = Comment(r[key], "разбор объявления")
        ws.row_dimensions[n + 1].height = 150

    last = len(rows) + 1
    ws.auto_filter.ref = f"A1:O{last}"
    ws.freeze_panes = "C2"

    note = (
        "Три дома одного типа: угловое здание, коммерция на первом этаже, квартиры выше. "
        "Признаки «угловой», «что внизу» и «что сверху» funda in business отдельными полями не даёт — "
        "они прочитаны из текста объявления и проверены по фотографиям фасада; в колонке «Угловой» "
        "написано, чем именно подтверждено. Площади, доход и доходность взяты из объявления там, где "
        "продавец их назвал: пустая ячейка значит «не сказано», а не «нет». Наведите курсор на ячейку "
        "с пометкой — в примечании написано, откуда число. «€/м²» считает сама книга формулой от цены "
        "и площади. Состояние — суждение по снимкам, а не поле funda."
    )
    c = ws.cell(last + 2, 1, note)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = Font(size=9, color="52606D")
    ws.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=15)
    ws.row_dimensions[last + 2].height = 96

    wb.save(a.out)
    print(f"{a.out}: {len(rows)} лотов, {len(COLS)} колонок")
    return 0


if __name__ == "__main__":
    sys.exit(main())
