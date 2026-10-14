#!/usr/bin/env python3
"""Лог осмотра квартир локальной моделью (VLM) в Excel: по каждой квартире дашборда — тексты и оценки модели.

python3 vlm_log_xlsx.py [база] → reports/Академический_оценки_квартир_моделью.xlsx
"""
import json
import math
import sqlite3
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from market_build import COND_ORDER, ROOMS_RU
from market_condition import READY
from market_vlm import AGE_RU, COND_RU, WORKS_RU, listing_note

DB = sys.argv[1] if len(sys.argv) > 1 else 'market-data/akademicheskiy.sqlite'
OUT = 'reports/Академический_оценки_квартир_моделью.xlsx'
GOLD = 'market-data/gold/gold.json'
FURN_RU = {'none': 'нет', 'partial': 'частично', 'full': 'полностью'}
REPAIR_RU = {'no': 'без ремонта', 'cosmetic': 'косметический', 'euro': 'евроремонт', 'design': 'дизайнерский'}
DECOR_RU = {'without': 'без отделки', 'preFine': 'предчистовая', 'fine': 'чистовая', 'fineWithFurniture': 'чистовая с мебелью'}
SELLER_MAP = {'no': 'требует ремонта', 'cosmetic': 'косметический', 'euro': 'современный', 'design': 'дизайнерский',
              'without': 'без отделки', 'preFine': 'white box', 'fine': 'современный', 'fineWithFurniture': 'современный'}

FONT = 'Arial'
HEAD_FILL = PatternFill('solid', fgColor='DDE3EA')
TEXT_FILL = PatternFill('solid', fgColor='FAF7EE')      # колонки с текстами модели
GRID = Side(style='thin', color='C8CDD3')
BORDER = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)

# (заголовок, ширина, пояснение, текст модели?)
COLS = [
    ('№', 5, None, False),
    ('ID ЦИАН', 12, 'Номер объявления, ссылка на карточку', False),
    ('Адрес', 30, None, False),
    ('Комнат', 9, None, False),
    ('Площадь, м²', 9, None, False),
    ('Этаж', 7, None, False),
    ('Год дома', 7, None, False),
    ('Цена, млн ₽', 9, 'Последняя цена объявления', False),
    ('Ремонт по объявлению', 14, 'Как ремонт указан продавцом в карточке ЦИАН', False),
    ('Вывод модели', 60, 'Поле summary — единственное поле, которое модель пишет по-русски (до 240 знаков)', True),
    ('Состояние', 14, 'condition: шкала состояния из inspector_prompt.md', False),
    ('Уверенность', 9, 'condition_confidence — самооценка модели', False),
    ('Возраст ремонта', 12, 'renovation_age', False),
    ('Нужные работы', 14, 'works', False),
    ('Комнаты: пол, стены, потолок', 34, 'rooms — заметка модели (англ., до 70 знаков)', True),
    ('Кухня', 34, 'kitchen — заметка модели (англ., до 70 знаков)', True),
    ('Санузел', 34, 'bathroom — заметка модели (англ., до 70 знаков)', True),
    ('Окна и двери', 30, 'windows_doors — заметка модели (англ., до 70 знаков)', True),
    ('Мебель', 10, 'furniture', False),
    ('Дефекты', 34, 'issues — до 4 пунктов по 40 знаков (англ.)', True),
    ('Особенности', 28, 'extras — до 3 пунктов по 30 знаков (англ.)', True),
    ('Фото в объявлении', 9, None, False),
    ('Фото показано модели', 9, 'После отбора: без дублей, планировок и фасадов, не больше 32', False),
    ('Фото с интерьером (по модели)', 10, 'photos_used — сколько фото модель сочла интерьерными', False),
    ('Что модели сказали о доме', 26, 'Факты объявления, переданные модели вместе с фото (англ.)', False),
    ('Время осмотра, с', 9, None, False),
    ('Дата осмотра', 11, None, False),
    ('Итог на дашборде', 14, 'Метка состояния, которую показывает дашборд', False),
    ('Откуда итог', 30, 'Осмотр модели или запасной классификатор CLIP (см. «Пояснения»)', False),
    ('Расходится с объявлением', 11, 'Да — итог и продавец по разные стороны границы «с ремонтом / под ремонт»', False),
    ('Ручная проверка', 40, 'Эталонная ручная разметка (14 квартир): состояние · возраст · работы — комментарий', False),
]
SUMMARY_COL = 10


def lst(x):
    if not x:
        return None
    return '; '.join(str(i) for i in x) if isinstance(x, list) else str(x)


def seller_text(r):
    a = REPAIR_RU.get(r['repair'] or '')
    b = DECOR_RU.get(r['decoration'] or '')
    return ', '.join(x for x in (a, b) if x) or None


def main():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    L = {r['cian_id']: r for r in con.execute('select * from listing where gone_seen is null')}
    insp = {r['cian_id']: r for r in con.execute('select * from inspect')}
    cond = {r['cian_id']: r for r in con.execute('select * from cond')}
    gold = {int(k): v for k, v in json.load(open(GOLD)).items() if not k.startswith('_')}

    rows = []
    for cid, x in L.items():
        i = insp.get(cid)
        d = json.loads(i['data'] or '{}') if i else {}
        c = cond.get(cid)
        n_ad = len(json.loads(x['photos'] or '[]'))
        stub = bool(i) and not (i['raw'] or '').strip()
        if not i:
            summary, how = 'Осмотра нет: у объявления нет фото квартиры.' if n_ad == 0 else \
                'Осмотра нет: фото ещё не были переданы модели.', None
        else:
            summary = d.get('summary')
            if stub:
                summary = f'{summary} (модель не вызывалась: после отбора фото интерьеров не осталось)'
        label = c['label'] if c else None
        if not c:
            how = 'нет оценки'
        elif c['model'] == 'vlm':
            how = 'осмотр модели'
        elif not i:
            how = f"классификатор CLIP: осмотра нет (вероятность {c['conf']:.0%})"
        elif i['condition'] == 'не показана':
            how = f"классификатор CLIP: модель не увидела квартиру (вероятность {c['conf']:.0%})"
        elif i['condition'] in READY and x['is_new'] and x['decoration'] in ('without', 'preFine'):
            how = (f"правило: новостройка без отделки — «ремонт» на фото это рендеры или лобби; "
                   f"далее CLIP (вероятность {c['conf']:.0%})")
        else:
            how = f"классификатор CLIP (вероятность {c['conf']:.0%})"
        sl = SELLER_MAP.get(x['repair'] or '') or SELLER_MAP.get(x['decoration'] or '')
        mism = bool(label and sl and label != sl and (label in READY) != (sl in READY))
        g = gold.get(cid)
        gtxt = None
        if g:
            gtxt = ' · '.join(filter(None, (COND_RU.get(g.get('condition')), AGE_RU.get(g.get('age')),
                                            WORKS_RU.get(g.get('works'))))) + (f" — {g['note']}" if g.get('note') else '')
        price = x['price_last']
        floor = f"{x['floor']}/{x['floors']}" if x['floor'] and x['floors'] else (x['floor'] or None)
        rows.append(dict(
            key=(COND_ORDER.index(i['condition']) if i and i['condition'] in COND_ORDER else
                 (90 if i else 99), x['address'] or ''),
            cid=cid, url=x['url'],
            vals=[None, cid, (x['address'] or '').replace('Академический, ', '', 1), ROOMS_RU.get(x['rooms'], x['rooms']),
                  x['area'], floor, x['build_year'], round(price / 1e6, 2) if price else None, seller_text(x),
                  summary, i['condition'] if i else None, d.get('condition_confidence'), i['age'] if i else None,
                  i['works'] if i else None, d.get('rooms'), d.get('kitchen'), d.get('bathroom'), d.get('windows_doors'),
                  FURN_RU.get(d.get('furniture'), d.get('furniture')), lst(d.get('issues')), lst(d.get('extras')),
                  n_ad, len(json.loads(i['photos'] or '[]')) if i else None, d.get('photos_used'),
                  listing_note(con, cid) if i and not stub else None, round(i['seconds']) if i and not stub and i['seconds'] else None,
                  i['date'] if i else None, label, how, 'да' if mism else None, gtxt],
            raw=(i['raw'] or '') if i else '', model=i['model'] if i else None))
    rows.sort(key=lambda r: r['key'])

    wb = Workbook()
    ws = wb.active
    ws.title = 'Оценки модели'
    hfont, font, link = Font(name=FONT, bold=True, size=9), Font(name=FONT, size=9), Font(name=FONT, size=9, color='1F4E9A', underline='single')
    for k, (h, w, note, _t) in enumerate(COLS, 1):
        cell = ws.cell(1, k, h)
        cell.font, cell.fill, cell.border = hfont, HEAD_FILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        if note:
            cell.comment = Comment(note, 'лог')
        ws.column_dimensions[get_column_letter(k)].width = w
    ws.row_dimensions[1].height = 42
    for n, r in enumerate(rows, 1):
        rr = n + 1
        r['vals'][0] = n
        lines = 1
        for k, v in enumerate(r['vals'], 1):
            cell = ws.cell(rr, k, v)
            cell.font, cell.border = font, BORDER
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            if COLS[k - 1][3]:
                cell.fill = TEXT_FILL
            if isinstance(v, str):
                lines = max(lines, sum(math.ceil(max(len(p), 1) / max(COLS[k - 1][1] * 1.15, 1)) for p in v.split('\n')))
        ws.cell(rr, 2).hyperlink = r['url']
        ws.cell(rr, 2).font = link
        ws.cell(rr, 5).number_format = '0.0'
        ws.cell(rr, 8).number_format = '0.00'
        ws.cell(rr, 12).number_format = '0%'
        ws.row_dimensions[rr].height = min(12 * lines + 3, 160)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{len(rows) + 1}'

    raw = wb.create_sheet('Ответ модели целиком')
    for k, (h, w) in enumerate([('ID ЦИАН', 12), ('Модель', 26), ('Ответ модели как есть (JSON)', 120)], 1):
        cell = raw.cell(1, k, h)
        cell.font, cell.fill, cell.border = hfont, HEAD_FILL, BORDER
        raw.column_dimensions[get_column_letter(k)].width = w
    for n, r in enumerate(rows, 2):
        txt = r['raw'] or ('— модель не вызывалась' if r['model'] else '— осмотра нет')
        for k, v in enumerate([r['cid'], r['model'], txt], 1):
            cell = raw.cell(n, k, v)
            cell.font, cell.border = font, BORDER
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        raw.cell(n, 1).hyperlink = r['url']
        raw.cell(n, 1).font = link
        raw.row_dimensions[n].height = min(12 * (txt.count('\n') + 1 + len(txt) // 140) + 3, 300)
    raw.freeze_panes = 'B2'

    models = sorted({r['model'] for r in rows if r['model']})
    dates = sorted({r['vals'][26] for r in rows if r['vals'][26]})
    n = len(rows) + 1
    info = wb.create_sheet('Пояснения')
    info.column_dimensions['A'].width = 34
    info.column_dimensions['B'].width = 100
    text = [
        ('Лог осмотра квартир локальной моделью', None),
        ('Район', 'Академический, все активные объявления дашборда'),
        ('Модель', ', '.join(models) + ' — работает локально (llama.cpp), смотрит все фото квартиры за один запрос'),
        ('Инструкция модели', 'inspector_prompt.md: шкала состояния, возраста ремонта и работ; схема ответа — vlm_schema.json'),
        ('Даты осмотра', ' — '.join(dates[:1] + dates[-1:]) if dates else '—'),
        ('Квартир в логе', f'=COUNTA(\'Оценки модели\'!B2:B{n})'),
        ('из них с осмотром модели', f'=COUNTA(\'Оценки модели\'!K2:K{n})'),
        (None, None),
        ('Как читать', None),
        ('Вывод модели', 'Главный текст — одно предложение по-русски: что модель увидела и что нужно делать.'),
        ('Заметки по комнатам, кухне, санузлу, окнам, дефекты, особенности',
         'Модель пишет их по-английски коротким стилем заметок — так задано в инструкции модели. '
         'Тексты приведены дословно, без правок и перевода.'),
        ('Обрезанные тексты', 'Длина полей ограничена схемой ответа (заметки — 70 знаков, дефекты — 40, особенности — 30, '
                              'вывод — 240), поэтому часть текстов обрывается на полуслове.'),
        ('Ответ модели целиком', 'Второй лист — ответ модели в исходном виде (JSON), из которого собраны колонки.'),
        ('Что модели сказали о доме', 'Вместе с фото модель получает год постройки и признак новостройки: ремонт не может '
                                     'быть старше дома.'),
        ('Итог на дашборде', 'Если модель увидела квартиру, дашборд показывает её оценку. Если квартира на фото не показана '
                             '(только рендеры, фасад, планировка) или это новостройка без отделки с фото лобби и рендеров — '
                             'метку ставит запасной классификатор по CLIP-признакам фото, обученный на разметке продавцов.'),
        ('Расходится с объявлением', '«да» — модель и продавец по разные стороны границы «с ремонтом» (современный, '
                                     'дизайнерский) / «под ремонт» (всё остальное).'),
        ('Ручная проверка', '14 квартир размечены вручную по всем фото — для сверки с моделью.'),
        ('Колонки', 'У заголовков есть примечания с пояснением поля.'),
        (None, None),
        ('Состояние по модели', 'квартир'),
    ]
    for k, (a, b) in enumerate(text, 1):
        info.cell(k, 1, a).font = Font(name=FONT, bold=True, size=12 if k == 1 else 10)
        cb = info.cell(k, 2, b)
        cb.font = Font(name=FONT, size=10, bold=a == 'Состояние по модели')
        cb.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
        info.cell(k, 1).alignment = Alignment(wrap_text=True, vertical='top')
    k = len(text)
    for lab in ['не показана'] + COND_ORDER:
        k += 1
        info.cell(k, 1, lab).font = Font(name=FONT, size=10)
        c = info.cell(k, 2, f'=COUNTIF(\'Оценки модели\'!K2:K{n},A{k})')
        c.font, c.alignment = Font(name=FONT, size=10), Alignment(horizontal='left')
    for r in (6, 7):
        info.cell(r, 2).alignment = Alignment(horizontal='left')
    wb.save(OUT)
    print(OUT, len(rows), 'квартир;', sum(1 for r in rows if r['model']), 'с осмотром')


if __name__ == '__main__':
    main()
