#!/usr/bin/env python3
"""Стромынка, 2 как участок для девелопера: остаточная стоимость земли при сносе здания и новом строительстве.

Данные: stromynka-data/newbuild.json (новостройки рядом, ЦИАН), land.json и buildings.json (участки с утверждённой
площадью застройки), цифры рынка земли из «Ведомостей» (октябрь 2026), себестоимость офисного строительства IBC/Stone.
python3 stromynka_dev.py → reports/Стромынка2_участок_для_девелопера.pdf, reports/Стромынка2_девелопер.xlsx
"""
import collections
import json
import os
import statistics as st
from html import escape

import stromynka_report as R
from stromynka_report import m

D = R.D
PDF = 'reports/Стромынка2_участок_для_девелопера.pdf'
XLSX = 'reports/Стромынка2_девелопер.xlsx'
HTML = 'reports/stromynka2_dev.html'
CURRENT = 383.0                    # рыночная стоимость при текущем использовании (здание с МФЦ), stromynka_report.py
PLOT, PLOT_ALL = 484, 484 + 3129   # участок :4 и вместе с общим участком :1863

# допущения модели: (ключ, подпись, значение, формат, пояснение)
INPUTS = [
    ('fp', 'Пятно застройки, м²', 441, '#,##0', 'Пятно нынешнего здания по ЕГРН; участок 484 м² — шире не построить'),
    ('price', 'Цена продажи квартир, ₽/м²', 650_000, '#,##0', 'Новостройки до 1 км: «Ренессанс» 627 тыс., «КОД Сокольники» 642 тыс., «СТОУН Сокольники» 554 тыс.; клубный «Артель» — 780 тыс.'),
    ('retail', 'Цена продажи 1 этажа (стрит-ритейл), ₽/м²', 450_000, '#,##0', 'Первая линия у метро'),
    ('eff', 'Доля продаваемой площади квартир', 0.75, '0%', 'Остальное — лестницы, лифты, холлы, стены'),
    ('eff1', 'Доля продаваемой площади 1 этажа', 0.85, '0%', None),
    ('cost', 'Строительство наземной части, ₽/м²', 200_000, '#,##0', 'Офисы высокого класса в Москве — 220 тыс. ₽/м² в 2026 г. (IBC Real Estate и Stone); жильё без отделки — чуть дешевле'),
    ('ucost', 'Строительство подземного этажа, ₽/м²', 250_000, '#,##0', None),
    ('under', 'Подземный этаж (технический), м²', 441, '#,##0', 'Паркинг на участке 484 м² не поместится: нет места для рампы'),
    ('demo', 'Снос нынешнего здания, ₽', 9_500_000, '#,##0', '≈ 8 тыс. ₽ на м² из 1 184,5 м²'),
    ('arch', 'Археологические раскопки, ₽', 20_000_000, '#,##0', 'Оценка: участок на территории культурного слоя XVIII в.; смету даёт Департамент культурного наследия'),
    ('vri', 'Плата городу за смену ВРИ и изменение ПЗЗ, ₽', 40_000_000, '#,##0', 'Оценка: ВРИ сейчас «бытовые услуги»; точную сумму считает ДГИ'),
    ('design', 'Проектирование и экспертиза, доля стройки', 0.07, '0%', None),
    ('nets', 'Подключение к сетям, доля стройки', 0.06, '0%', None),
    ('mkt', 'Маркетинг и продажи, доля выручки', 0.04, '0%', None),
    ('profit', 'Прибыль девелопера, доля выручки', 0.20, '0%', 'Обычный ориентир 15–25%; маленький проект в охранной зоне — риск выше'),
    ('fin', 'Финансирование стройки, доля затрат', 0.09, '0%', 'Проектное финансирование с эскроу, ≈ 2 года'),
    ('bridge', 'Ставка по кредиту на покупку земли, годовых', 0.20, '0%', '«Ведомости»: бридж-кредиты на землю — более 20% годовых'),
    ('years', 'Срок согласований до начала стройки, лет', 2.5, '0.0', 'Смена ВРИ и ПЗЗ через ГЗК, согласование с ДКН, ГПЗУ, экспертиза, раскопки'),
    ('rent', 'Чистый доход от МФЦ, ₽ в год', 36_671_205, '#,##0', '38,95 млн с НДС минус эксплуатация 2,28 млн'),
    ('rent_years', 'Сколько лет МФЦ остаётся, пока идут согласования', 0, '0', '0 — здание сразу свободно; 2 — МФЦ платит, пока девелопер получает разрешения'),
    ('offer', 'Предложение покупателя, ₽', 550_000_000, '#,##0', None),
]
P0 = {k: v for k, _, v, _, _ in INPUTS}
FLOORS = [2, 3, 4, 5, 6, 8, 10]
OPTIMISTIC = dict(price=750_000, cost=180_000, profit=0.15, bridge=0.15)
# участки и здания под редевелопмент на ЦИАН с известной разрешённой площадью (из описаний): id → (м² ГНС, что разрешено)
GNS = {
    '328295784': (27_300, 'пройдена ГЗК, деловое управление, 17 этажей'),
    '332186074': (48_900, 'пройдена ГЗК, ТЭП утверждены; есть здание 10,8 тыс. м²'),
    '334643903': (33_140, 'ТЭП согласованы, участок выведен из КРТ'),
    '332610492': (12_800, 'ГПЗУ 2013 г. на 12,8 тыс. м² наземной части'),
    '335014595': (4_000, 'ГПЗУ 2025 г., нежилое, здание под снос'),
    '334567721': (3_000, 'плотность 15 тыс. м²/га, высота до 12 м'),
    '285429175': (1_146, 'ГПЗУ: до 1 146 м², высота до 65 м'),
    '333706127': (3_000, 'ГПЗУ на реконструкцию до 3 000 м², 7 этажей'),
    '334495464': (6_779, 'ГПЗУ на 6 778,8 м²; есть здание 3,2 тыс. м²'),
}
NB_NEAR = 2.3      # новостройки в этом радиусе от объекта, км


def model(f, p=None):
    p = {**P0, **(p or {})}
    above = p['fp'] * f
    resi = p['fp'] * (f - 1) * p['eff']
    first = p['fp'] * p['eff1']
    gdv = resi * p['price'] + first * p['retail']
    cons = above * p['cost'] + p['under'] * p['ucost']
    soft = cons * (p['design'] + p['nets'])
    fixed = p['demo'] + p['arch'] + p['vri']
    fin = (cons + soft) * p['fin']
    mkt, prof = gdv * p['mkt'], gdv * p['profit']
    res0 = gdv - cons - soft - fixed - fin - mkt - prof
    land = res0 / (1 + p['bridge']) ** p['years']
    n = p['rent_years']
    rent = p['rent'] * (1 - (1 + p['bridge']) ** -n) / p['bridge'] if n else 0
    return dict(f=f, above=above, resi=resi, first=first, gdv=gdv, cons=cons, soft=soft, fixed=fixed, fin=fin, mkt=mkt,
                prof=prof, res0=res0, land=land, rent=rent, total=land + rent, gns=above + p['under'],
                per_gns=res0 / (above + p['under']), cost_all=cons + soft + fixed + fin)


def floors_for(target, p=None):
    """Сколько наземных этажей нужно, чтобы девелопер мог заплатить target рублей (модель линейна по этажам)."""
    a, b = model(10, p), model(11, p)
    step = b['total'] - a['total']
    return 10 + (target - a['total']) / step


def newbuild():
    N = json.load(open(f'{D}/newbuild.json'))
    g = collections.defaultdict(list)
    for x in N.values():
        if x['price'] and x['area'] and x['km_obj'] is not None:
            g[x['jk']].append(x)
    rows = []
    for jk, xs in g.items():
        km = st.median(x['km_obj'] for x in xs)
        if km > NB_NEAR or len(xs) < 3:
            continue
        pp = [x['price'] / x['area'] for x in xs]
        mt = collections.Counter((x['metro'][0][0], x['metro'][0][1]) for x in xs if x['metro'] and x['metro'][0][2] == 'walk')
        fin = collections.Counter((x['finish'] or {}).get('year') for x in xs).most_common(1)[0][0]
        rows.append(dict(jk=jk, km=km, n=len(xs), med=st.median(pp), lo=min(pp), hi=max(pp), apart=sum(bool(x['apart']) for x in xs),
                         metro=mt.most_common(1)[0][0] if mt else None, finish=fin, floors=xs[0]['floors'],
                         url=next((x['jk_url'] for x in xs if x.get('jk_url')), xs[0]['url'])))
    return sorted(rows, key=lambda r: r['km'])


def gns_comps():
    A = {**json.load(open(f'{D}/land.json')), **json.load(open(f'{D}/buildings.json'))}
    rows = []
    for cid, (gns, note) in GNS.items():
        x = A[cid]
        land = {'332186074': 9000, '334643903': 21184}.get(cid, x['land_m2'])      # площадь по описанию, где ЦИАН путает
        rows.append(dict(id=cid, url=x['url'].split('?')[0], raion=x['raion'], address=x['address'], km=x['km_kremlin'], land=land,
                         price=x['price'], gns=gns, per=x['price'] / gns, note=note, vat=x['vat']))
    return sorted(rows, key=lambda r: r['per'])


# ---------- графика ----------

def curve_png(out):
    matplotlib, plt = R.plt_setup()
    fs = [f / 2 for f in range(4, 41)]
    base = [model(f)['total'] / 1e6 for f in fs]
    rent = [model(f, dict(rent_years=2))['total'] / 1e6 for f in fs]
    opt = [model(f, OPTIMISTIC)['total'] / 1e6 for f in fs]
    fig, ax = plt.subplots(figsize=(7.4, 3.6), dpi=200)
    fig.patch.set_facecolor(R.SURF)
    ax.set_facecolor(R.SURF)
    ax.axvspan(3, 6, color='#e9e7e0', zorder=0)
    ax.text(4.5, -140, 'реально в охранной\nзоне: 3–6 этажей', ha='center', va='bottom', fontsize=8.5, color=R.INK2)
    ax.plot(fs, opt, color=R.AQUA, lw=1.6, ls='--', label='оптимистично: квартиры 750 тыс./м², стройка 180 тыс., прибыль 15%, кредит 15%')
    ax.plot(fs, rent, color=R.BLUE, lw=1.2, alpha=0.55, label='база + МФЦ платит 2 года, пока идут согласования')
    ax.plot(fs, base, color=R.BLUE, lw=2.2, label='база: квартиры 650 тыс./м², стройка 200 тыс., прибыль 20%, кредит 20%')
    for v, lab, col in ((P0['offer'] / 1e6, 'предложение 550 млн', R.ORANGE), (CURRENT, 'здание с МФЦ сейчас ≈ 383 млн', R.INK)):
        ax.axhline(v, color=col, lw=1.4, ls='-' if col == R.ORANGE else ':')
        ax.text(2.15, v + 9, lab, ha='left', fontsize=8.5, color=col)
    ax.axhline(0, color=R.INK2, lw=0.8)
    ax.set_xlim(2, 20)
    ax.set_ylim(-150, 650)
    ax.set_xticks(range(2, 21, 2))
    ax.set_xlabel('Этажей над землёй на пятне 441 м²', color=R.INK2)
    ax.set_ylabel('Сколько девелопер может\nзаплатить за участок, млн руб.', color=R.INK2)
    ax.grid(color=R.GRID, lw=0.8)
    ax.set_axisbelow(True)
    for s in ('top', 'right'):
        ax.spines[s].set_visible(False)
    for s in ('left', 'bottom'):
        ax.spines[s].set_color(R.GRID)
    ax.tick_params(colors=R.INK2)
    ax.legend(loc='lower right', bbox_to_anchor=(1.0, 0.04), frameon=False, fontsize=7.6, labelcolor=R.INK2)
    fig.tight_layout()
    fig.savefig(out, facecolor=R.SURF)
    plt.close(fig)


def ha_png(out):
    matplotlib, plt = R.plt_setup()
    rows = [('Молжаниново: 5,16 га за 7,19 млрд', 7.19 / 5.16, R.GRID),
            ('Средняя цена 1 га в Москве (IBC)', 3.9, R.GRID),
            ('Павелецкий вокзал: 0,47 га за 5,5 млрд', 5.5 / 0.47, R.GRID),
            ('Стромынка, 2 за 550 млн: участок 484 м²', 550 / 1e3 / (PLOT / 1e4), R.ORANGE),
            ('То же, если считать и общий участок 3 129 м²', 550 / 1e3 / (PLOT_ALL / 1e4), '#f4b393')]
    fig, ax = plt.subplots(figsize=(7.4, 2.0), dpi=200)
    fig.patch.set_facecolor(R.SURF)
    ax.set_facecolor(R.SURF)
    ys = list(range(len(rows)))[::-1]
    ax.barh(ys, [r[1] for r in rows], color=[r[2] if r[2] != R.GRID else '#9db8d9' for r in rows], height=0.6, zorder=2)
    ax.set_yticks(ys)
    ax.set_yticklabels([r[0] for r in rows], fontsize=8.5, color=R.INK2)
    for y, r in zip(ys, rows):
        ax.text(r[1] + 0.15, y, m(r[1], 1), va='center', fontsize=8.5, color=R.INK2)
    ax.set_xlabel('Млрд ₽ за гектар'.replace('₽', 'руб.'), color=R.INK2)
    ax.set_xlim(0, 13.5)
    ax.grid(axis='x', color=R.GRID, lw=0.8)
    ax.set_axisbelow(True)
    for s in ('top', 'right', 'left'):
        ax.spines[s].set_visible(False)
    ax.spines['bottom'].set_color(R.GRID)
    ax.tick_params(colors=R.INK2, length=0)
    fig.tight_layout()
    fig.savefig(out, facecolor=R.SURF)
    plt.close(fig)


def nb_png(nb, out):
    matplotlib, plt = R.plt_setup()
    rows = sorted(nb, key=lambda r: r['med'])
    fig, ax = plt.subplots(figsize=(7.4, 2.1), dpi=200)
    fig.patch.set_facecolor(R.SURF)
    ax.set_facecolor(R.SURF)
    ys = range(len(rows))
    ax.barh(list(ys), [r['med'] / 1e3 for r in rows], color=[R.ORANGE if r['apart'] > r['n'] / 2 else R.BLUE for r in rows],
            height=0.6, zorder=2)
    ax.set_yticks(list(ys))
    ax.set_yticklabels([f"{r['jk']} · {m(r['km'], 1)} км" for r in rows], fontsize=8.5, color=R.INK2)
    for y, r in zip(ys, rows):
        ax.text(r['med'] / 1e3 + 6, y, m(r['med'] / 1e3), va='center', fontsize=8, color=R.INK2)
    v = P0['price'] / 1e3
    ax.axvline(v, color=R.INK, lw=1.4, ls='--')
    ax.text(v + 4, len(rows) - 0.45, f'в расчёте: {m(v)} тыс./м²', fontsize=8.5, color=R.INK)
    ax.set_xlabel('Медиана цены квартир в продаже, тыс. руб. за м²', color=R.INK2)
    ax.set_xlim(0, 880)
    ax.set_ylim(-0.6, len(rows) + 0.2)
    ax.grid(axis='x', color=R.GRID, lw=0.8)
    ax.set_axisbelow(True)
    for s in ('top', 'right', 'left'):
        ax.spines[s].set_visible(False)
    ax.spines['bottom'].set_color(R.GRID)
    ax.tick_params(colors=R.INK2, length=0)
    fig.tight_layout()
    fig.savefig(out, facecolor=R.SURF)
    plt.close(fig)


# ---------- Excel ----------

def write_xlsx(nb, gc):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    F = 'Arial'
    bold, font = Font(name=F, bold=True, size=9), Font(name=F, size=9)
    blue, note = Font(name=F, size=9, color='0000FF'), Font(name=F, size=8.5, color='555555')
    grid = Side(style='thin', color='C8CDD3')
    border = Border(left=grid, right=grid, top=grid, bottom=grid)
    head = PatternFill('solid', fgColor='DDE3EA')
    yellow = PatternFill('solid', fgColor='FFF2CC')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Модель девелопера'
    ws.column_dimensions['A'].width = 52
    for c in range(2, 10):
        ws.column_dimensions[L(c)].width = 13
    ws.column_dimensions['J'].width = 70
    ws['A1'] = 'Стромынка, 2: сколько девелопер может заплатить за участок (снос здания, новое строительство)'
    ws['A1'].font = Font(name=F, bold=True, size=11)
    ws['A2'] = 'Синий шрифт — допущения, их можно менять; всё остальное считается формулами. Жёлтые — главные допущения.'
    ws['A2'].font = note
    ref = {}
    for i, (k, lab, v, fmt, why) in enumerate(INPUTS, 4):
        ws.cell(i, 1, lab).font = font
        c = ws.cell(i, 2, v)
        c.font, c.number_format = blue, fmt
        if k in ('price', 'cost', 'profit', 'bridge', 'rent_years'):
            c.fill = yellow
        if why:
            ws.cell(i, 3, why).font = note
        ref[k] = f'$B${i}'
    r0 = 4 + len(INPUTS) + 1
    ws.cell(r0, 1, 'РАСЧЁТ ПО ВАРИАНТАМ ЭТАЖНОСТИ').font = bold
    lines = [
        ('Этажей над землёй', None, '0'),
        ('Наземная площадь, м²', '={fp}*{c}{r_f}', '#,##0'),
        ('Продаваемые квартиры (2 этаж и выше), м²', '={fp}*({c}{r_f}-1)*{eff}', '#,##0'),
        ('Продаваемый 1 этаж, м²', '={fp}*{eff1}', '#,##0'),
        ('Выручка, ₽', '={c}{r_resi}*{price}+{c}{r_first}*{retail}', '#,##0'),
        ('Строительство, ₽', '={c}{r_above}*{cost}+{under}*{ucost}', '#,##0'),
        ('Проектирование и сети, ₽', '={c}{r_cons}*({design}+{nets})', '#,##0'),
        ('Снос, раскопки, плата за ВРИ и ПЗЗ, ₽', '={demo}+{arch}+{vri}', '#,##0'),
        ('Финансирование стройки, ₽', '=({c}{r_cons}+{c}{r_soft})*{fin}', '#,##0'),
        ('Маркетинг и продажи, ₽', '={c}{r_gdv}*{mkt}', '#,##0'),
        ('Прибыль девелопера, ₽', '={c}{r_gdv}*{profit}', '#,##0'),
        ('Остаётся на землю к началу стройки, ₽', '={c}{r_gdv}-SUM({c}{r_cons}:{c}{r_prof})', '#,##0'),
        ('Стоимость земли сегодня (с учётом кредита на срок согласований), ₽', '={c}{r_res0}/(1+{bridge})^{years}', '#,##0'),
        ('Доход от МФЦ на время согласований (сегодняшние деньги), ₽', '=IF({rent_years}=0,0,{rent}*(1-(1+{bridge})^(-{rent_years}))/{bridge})', '#,##0'),
        ('ИТОГО девелопер может заплатить, млн ₽', '=({c}{r_land}+{c}{r_rentpv})/1e6', '#,##0.0'),
        ('Предложение 550 млн больше этой суммы, раз', '=IF({c}{r_total}<=0,"—",{offer}/1e6/{c}{r_total})', '0.0'),
        ('Остаток на землю на 1 м² общей площади (наземной и подземной), ₽', '={c}{r_res0}/({c}{r_above}+{under})', '#,##0'),
    ]
    keys = ['r_' + k for k in ('f', 'above', 'resi', 'first', 'gdv', 'cons', 'soft', 'fixed', 'fin', 'mkt', 'prof', 'res0', 'land',
                               'rentpv', 'total', 'x', 'pergns')]
    rows = {k: r0 + 1 + i for i, k in enumerate(keys)}
    for k, (lab, _, _) in zip(keys, lines):
        ws.cell(rows[k], 1, lab).font = bold if k in ('r_total', 'r_res0') else font
    for j, fl in enumerate(FLOORS, 2):
        c = L(j)
        for k, (lab, frm, fmt) in zip(keys, lines):
            if k == 'r_f':
                cell = ws.cell(rows[k], j, fl)
                cell.font = blue
            else:
                cell = ws.cell(rows[k], j, frm.format(c=c, **ref, **rows))
                cell.font = Font(name=F, size=9, bold=k in ('r_total', 'r_res0'))
            cell.number_format = fmt
    rn = rows['r_pergns'] + 2
    ws.cell(rn, 1, 'Сколько этажей нужно, чтобы девелопер окупил 550 млн').font = bold
    a, b = f'{L(2 + FLOORS.index(10))}{rows["r_total"]}', f'{L(2 + FLOORS.index(8))}{rows["r_total"]}'
    ws.cell(rn, 2, f'=10+({ref["offer"]}/1e6-{a})/(({a}-{b})/2)').number_format = '0.0'
    ws.cell(rn, 2).font = bold
    ws.cell(rn, 3, 'Модель линейна по этажам: каждый этаж добавляет одинаковую сумму').font = note
    ws.cell(rn + 1, 1, 'Здание с МФЦ при текущем использовании, млн ₽').font = font
    ws.cell(rn + 1, 2, CURRENT).font = blue
    ws.cell(rn + 1, 3, 'Доходный и сравнительный подходы, отчёт «Стромынка2_оценка.pdf»').font = note
    ws.freeze_panes = 'B4'

    # ориентиры рынка земли
    s2 = wb.create_sheet('Ориентиры рынка земли')
    s2.column_dimensions['A'].width = 62
    s2.column_dimensions['B'].width = 16
    s2.column_dimensions['C'].width = 70
    bench = [
        ('Площадь участка :4, га', PLOT / 1e4, '0.0000', 'ЕГРН', True),
        ('Средняя цена 1 га под жильё в Москве, 9 мес. 2026, млрд ₽', 3.9, '0.0', 'IBC Real Estate, «Ведомости»', True),
        ('Минимум 2026: Молжаниново, млрд ₽ за га', 1.4, '0.0', 'Asseto, «Ведомости»: 5,16 га за 7,19 млрд', True),
        ('Максимум 2026: у Павелецкого вокзала, млрд ₽ за га', 12, '0.0', 'Asseto, «Ведомости»: 0,47 га за 5,5 млрд', True),
        ('Участок по средней цене, млн ₽', '=B2*B3*1000', '#,##0', None, False),
        ('Участок по минимальной цене, млн ₽', '=B2*B4*1000', '#,##0', None, False),
        ('Участок по максимальной цене, млн ₽', '=B2*B5*1000', '#,##0', None, False),
        ('Предложение 550 млн в пересчёте на 1 га, млрд ₽', f"='Модель девелопера'!{ref['offer']}/1e9/B2", '0.0', None, False),
        ('Доля земли в затратах жилого проекта без документов', 0.125, '0.0%', '«Ведомости»: 12–13%', True),
        ('Доля земли с готовой разрешительной документацией', 1 / 3, '0.0%', '«Ведомости»: до трети', True),
        ('Затраты проекта без земли, 6 этажей, млн ₽', None, '#,##0', 'Стройка, проектирование, сети, снос, раскопки, ВРИ, финансирование — из модели', False),
        ('Земля по доле 12,5% (6 этажей), млн ₽', '=B12*B10/(1-B10)', '#,##0', None, False),
        ('Земля по доле 1/3 — если бы ГПЗУ уже был (6 этажей), млн ₽', '=B12*B11/(1-B11)', '#,##0', None, False),
        ('Медиана цены участков на ЦИАН за 1 м² разрешённой площади, ₽', round(st.median(r['per'] for r in gc)), '#,##0', 'Лист «Участки с ГНС (ЦИАН)», нежилое', True),
        ('Земля по этой цене, 6 этажей + подвал, млн ₽', None, '#,##0', None, False),
    ]
    c6 = L(2 + FLOORS.index(6))
    for i, (lab, v, fmt, src, inp) in enumerate(bench, 2):
        s2.cell(i, 1, lab).font = font
        if lab.startswith('Затраты проекта без земли'):
            v = f"=('Модель девелопера'!{c6}{rows['r_cons']}+'Модель девелопера'!{c6}{rows['r_soft']}+'Модель девелопера'!{c6}{rows['r_fixed']}+'Модель девелопера'!{c6}{rows['r_fin']})/1e6"
        if lab.startswith('Земля по этой цене'):
            v = f"=B15*('Модель девелопера'!{c6}{rows['r_above']}+'Модель девелопера'!B{4 + [k for k, *_ in INPUTS].index('under')})/1e6"
        c = s2.cell(i, 2, v)
        c.number_format = fmt
        c.font = blue if inp else Font(name=F, size=9, color='008000' if isinstance(v, str) and "'" in v else '000000')
        if src:
            s2.cell(i, 3, src).font = note

    def table(title, cols, data, links=(1,)):
        sh = wb.create_sheet(title)
        for k, (h, w) in enumerate(cols, 1):
            c = sh.cell(1, k, h)
            c.font, c.fill, c.border = bold, head, border
            c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
            sh.column_dimensions[L(k)].width = w
        sh.row_dimensions[1].height = 32
        for n, (vals, url) in enumerate(data, 2):
            for k, v in enumerate(vals, 1):
                c = sh.cell(n, k, v)
                c.font, c.border = font, border
                c.alignment = Alignment(wrap_text=True, vertical='top')
            if url:
                for k in links:
                    sh.cell(n, k).hyperlink = url
                    sh.cell(n, k).font = Font(name=F, size=9, color='1F4E9A', underline='single')
        sh.freeze_panes = 'B2'
        return sh
    table('Участки с ГНС (ЦИАН)', [('ID ЦИАН', 11), ('Район', 14), ('Адрес', 26), ('До центра, км', 8), ('Участок, м²', 10),
                                    ('Цена, млн ₽', 10), ('Разрешённая площадь, м²', 11), ('₽ за м² разрешённой площади', 12),
                                    ('НДС', 11), ('Что разрешено', 50)],
          [([r['id'], r['raion'], r['address'], r['km'], round(r['land']), round(r['price'] / 1e6), r['gns'], round(r['per']),
             r['vat'], r['note']], r['url']) for r in gc])
    table('Новостройки рядом (ЦИАН)', [('ЖК', 26), ('От объекта, км', 9), ('Метро, пешком', 22), ('Срок сдачи', 9), ('Этажей', 8),
                                       ('Лотов', 7), ('Медиана ₽/м²', 12), ('Мин ₽/м²', 11), ('Макс ₽/м²', 11), ('Апартаменты', 10)],
          [([r['jk'], r['km'], f"{r['metro'][0]}, {r['metro'][1]} мин" if r['metro'] else None, r['finish'], r['floors'], r['n'],
             round(r['med']), round(r['lo']), round(r['hi']), r['apart']], r['url']) for r in nb])
    s5 = wb.create_sheet('Статья «Ведомостей»')
    s5.column_dimensions['A'].width = 120
    for i, t in enumerate(ARTICLE, 1):
        s5.cell(i, 1, t).font = bold if i == 1 else font
        s5.cell(i, 1).alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(XLSX)
    return rows


ARTICLE = [
    '«Девелоперы запасаются землёй», «Ведомости», октябрь 2026 г. (Мария Асиновская) — главное',
    'Ricci: за 9 мес. 2026 г. девелоперы вложили в участки под жильё 211 млрд ₽ (+37%), 92 сделки (+35%); на столичный регион — около 179 млрд ₽, 85%.',
    'Nikoliers: 251,5 млрд ₽ (+27%), площадь купленных участков −10%, до 952 га; Москва с Подмосковьем — 170,6 млрд ₽ (+16%), площадь −54%, до 419 га.',
    'IBC Real Estate: 161 млрд ₽ за три квартала (+15%); в столице — 95 млрд ₽ против 96 млрд годом ранее; в январе–сентябре закрыто 29 сделок против 34.',
    'NF Group: Московский регион, Петербург и Ленобласть — 123,5 млрд ₽ (−13,6%), сделок 41 против 56.',
    'IBC Real Estate: средняя стоимость 1 га в Москве за 9 месяцев — 3,9 млрд ₽, на 10% выше прошлого года. Asseto: от 1,4 млрд ₽ за га (Молжаниново, 5,16 га за 7,19 млрд) до 12 млрд ₽ (0,47 га у Павелецкого вокзала за 5,5 млрд); средняя цена за год +16%.',
    'Ricci: более 70% сделок — участки с одобренными параметрами застройки; готовые документы сокращают сроки и риски.',
    'Доля участка в затратах на жилой проект — 12–13%, с оформленной разрешительной документацией — до трети (Asseto).',
    'Бридж-кредиты на покупку земли — более 20% годовых (Element); девелоперы рассчитывают на снижение ключевой ставки.',
    'Девелоперский цикл — 5–8 лет (BMS): покупая участок сейчас, застройщики ориентируются на 2030-е годы.',
    'В Москве в приоритете бизнес-класс и выше — их покупатели меньше зависят от ипотечных ставок.',
    'Продавцы стали сговорчивее; на рынок выходят участки, которые раньше не продавались; хорошие участки не дешевеют, но договориться проще.',
    'Для сравнения: в первом полугодии 2026 г. средняя цена 1 га в Москве была 3,6 млрд ₽ (−12,2% за год) — «Ведомости», 14.07.2026.',
]


# ---------- PDF ----------

def build_html(nb, gc):
    S = lambda v: escape(str(v)) if v is not None else '—'
    base = {f: model(f) for f in FLOORS}
    rent2 = {f: model(f, dict(rent_years=2)) for f in FLOORS}
    opt6 = model(6, OPTIMISTIC)
    need = floors_for(P0['offer'])
    need_rent = floors_for(P0['offer'], dict(rent_years=2))
    need_opt = floors_for(P0['offer'], OPTIMISTIC)
    ha_offer = P0['offer'] / 1e9 / (PLOT / 1e4)
    gmed = st.median(r['per'] for r in gc)
    g4, g6 = base[4]['gns'], base[6]['gns']
    share = lambda f, s: base[f]['cost_all'] * s / (1 - s) / 1e6
    lo, hi = 50, 200
    rows = ''
    n = lambda v, d=0: m(v, d).replace('-', '−')
    for f in FLOORS:
        b, r2 = base[f], rent2[f]
        rows += (f'<tr class="{"hl" if f == 6 else ""}"><td>{f}{" (как сейчас)" if f == 2 else ""}</td><td class=r>{n(b["above"])}</td>'
                 f'<td class=r>{n(b["resi"] + b["first"])}</td><td class=r>{n(b["gdv"] / 1e6)}</td>'
                 f'<td class=r>{n((b["cons"] + b["soft"] + b["fixed"] + b["fin"]) / 1e6)}</td><td class=r>{n((b["mkt"] + b["prof"]) / 1e6)}</td>'
                 f'<td class=r>{n(b["res0"] / 1e6)}</td><td class=r><b>{n(b["total"] / 1e6)}</b></td><td class=r>{n(r2["total"] / 1e6)}</td></tr>')
    nb_rows = ''.join(f'<tr><td><a href="{r["url"]}">{escape(r["jk"])}</a></td><td class=r>{m(r["km"], 1)}</td>'
                      f'<td>{S(r["metro"][0]) + ", " + S(r["metro"][1]) + " мин" if r["metro"] else "—"}</td><td class=r>{S(r["finish"])}</td>'
                      f'<td class=r>{S(r["floors"])}</td><td class=r>{r["n"]}</td><td class=r>{m(r["med"] / 1e3)}</td></tr>' for r in nb)
    g_rows = ''.join(f'<tr><td><a href="{r["url"]}">{r["id"]}</a></td><td>{S(r["raion"])}</td><td class=r>{m(r["land"])}</td>'
                     f'<td class=r>{m(r["price"] / 1e6)}</td><td class=r>{m(r["gns"])}</td><td class=r>{m(r["per"] / 1e3)}</td>'
                     f'<td>{escape(r["note"])}</td></tr>' for r in gc)
    methods = [
        ('Расчёт от будущего проекта, 4–6 этажей, база', f'{n(base[4]["total"] / 1e6)}…{n(base[6]["total"] / 1e6)}'),
        ('То же, если МФЦ платит аренду 2 года, пока идут согласования', f'{m(rent2[4]["total"] / 1e6)}…{m(rent2[6]["total"] / 1e6)}'),
        ('То же, оптимистично (6 этажей)', f'{m(opt6["total"] / 1e6)}'),
        ('Доля земли 12–13% затрат проекта («Ведомости»), 4–6 этажей', f'{m(share(4, 0.125))}…{m(share(6, 0.125))}'),
        ('Участки с разрешённой площадью на ЦИАН (медиана {} тыс. ₽ за м²), 4–6 этажей'.format(m(gmed / 1e3)), f'{m(gmed * g4 / 1e6)}…{m(gmed * g6 / 1e6)}'),
        ('Средняя цена 1 га в Москве — 3,9 млрд ₽ (IBC)', m(3.9e9 * PLOT / 1e4 / 1e6)),
        ('Для сравнения: если бы ГПЗУ на 4–6 этажей уже был (доля земли — треть затрат)', f'{m(share(4, 1 / 3))}…{m(share(6, 1 / 3))}'),
    ]
    meth_rows = ''.join(f'<tr class="{"muted" if i == len(methods) - 1 else ""}"><td>{escape(a)}</td><td class=r>{b} млн</td></tr>'
                        for i, (a, b) in enumerate(methods))
    css = R.CSS + '''
.t tr.muted td { color: #555; font-style: italic; }
.kpis { grid-template-columns: 1.1fr 1fr 1fr 1fr; }
'''
    body = f'''
<div class=t1>Стромынка, 2 как участок для девелопера</div>
<div class=t2>Сколько стоит земля, если здание снести и строить заново · данные ЦИАН и рынка земли на 19 октября 2026 г.</div>
<div class=kpis>
<div class=kpi><div class=v>{lo}–{hi} млн</div><div class=l>стоимость участка для девелопера<br>(даже оптимистично — до {m(round(opt6["total"] / 1e7) * 10)} млн)</div></div>
<div class=kpi><div class=v>{m(CURRENT)} млн</div><div class=l>здание с МФЦ при текущем<br>использовании — дороже земли</div></div>
<div class=kpi><div class=v>{m(ha_offer, 1)} млрд/га</div><div class=l>550 млн в пересчёте на гектар;<br>в среднем по Москве — 3,9</div></div>
<div class=kpi><div class=v>≈ {m(need)} этажей</div><div class=l>нужно построить на пятне 441 м²,<br>чтобы 550 млн окупились</div></div>
</div>
<div class=box>
<p><b>Для девелопера наш участок стоит около {lo}–{hi} млн.</b> 550 млн — в {m(550 / hi, 1)} раза больше верхней границы и в
{m(550 / (opt6["total"] / 1e6), 1)} раза больше самого оптимистичного расчёта (≈ {m(round(opt6["total"] / 1e7) * 10)} млн). Участок маленький (484 м²), по нему
проходят охранная зона объекта культурного наследия, зона регулирования застройки и культурный слой XVIII века. Разрешённое использование —
«бытовые услуги», под жильё его нужно менять через градостроительно-земельную комиссию. По нашей оценке, реально здесь можно построить дом на 3–6 этажей.
Чтобы заплатить 550 млн и заработать обычную для девелопера прибыль, нужен дом примерно на {m(need)} этажей
(при оптимистичных допущениях — на {m(need_opt)}). В охранной зоне это практически невозможно.</p>
<p><b>Земля под зданием стоит меньше, чем здание с МФЦ.</b> Сносить его девелоперу невыгодно: текущее использование даёт около {m(CURRENT)} млн,
а проект на месте здания — {lo}–{hi} млн. Поэтому 550 млн нельзя объяснить ценой земли. Скорее всего, у покупателя есть свой интерес: соседние здания
и общий участок :1863 (3 129 м²), сведения о будущем проекте на квартале или планы на само здание. Возможно и другое: он рассчитывает снизить
цену после проверки документов.</p>
<p><b>Что делать:</b> выяснить, кто покупатель и связан ли он с соседними зданиями :1735 и :4604 на общем участке. Брать задаток и не соглашаться
на сделку с условием «если получу разрешения»: риск согласований должен остаться на покупателе. До сделки не продлевать договор с МФЦ на 3 года без
права досрочного расторжения.</p></div>

<h1>1. Что можно построить на участке</h1>
<table class=t><tbody>
<tr><td style="width:27%">Участок</td><td>:4 — 484 м², пятно здания 441 м², аренда у города до 2066 г.; рядом общий с соседями участок :1863 (3 129 м²)
с ещё двумя зданиями — его девелопер без соседей не получит</td></tr>
<tr><td>Разрешённое использование</td><td>«Эксплуатация здания под бытовые услуги». Для жилья или апартаментов нужно сменить ВРИ и, скорее всего,
поправить правила землепользования и застройки (решение ГЗК), заплатить городу за смену ВРИ и оформить ГПЗУ</td></tr>
<tr><td>Охранная зона ОКН № 186</td><td>По закону № 73-ФЗ (ст. 34) в охранной зоне строительство запрещено, кроме мер по регенерации
исторической среды: объём и высоту нового здания задаёт режим зоны, проект согласует Департамент культурного наследия</td></tr>
<tr><td>Зона регулирования застройки № 1</td><td>Ограничивает высоту и архитектурные решения</td></tr>
<tr><td>Культурный слой XVIII в.</td><td>До начала земляных работ — археологические раскопки за счёт застройщика, по сезону (июнь–октябрь)</td></tr>
<tr><td>Подземный паркинг</td><td>На 484 м² не помещается рампа — машиномест почти не будет, для премиальных квартир это минус</td></tr>
</tbody></table>
<p class=src>По выпискам ЕГРН от 05.10.2026. Точные параметры (высота, плотность) покажет ГПЗУ — его можно бесплатно получить через mos.ru за 14 рабочих дней.</p>

<section class=pb><h1>2. За сколько продавать квартиры</h1>
<img class=full src="file://{os.path.abspath(f"{D}/dev_nb.png")}">
<table class=t style="margin-top:6px"><thead><tr><th>ЖК</th><th class=r>От нас, км</th><th>Метро пешком</th><th class=r>Сдача</th><th class=r>Этажей</th>
<th class=r>Лотов</th><th class=r>Медиана, тыс. ₽/м²</th></tr></thead><tbody>{nb_rows}</tbody></table>
<p class=src>Новостройки в продаже на ЦИАН в радиусе {m(NB_NEAR, 1)} км, медиана цены по ЖК; на графике оранжевым — апартаменты.
Для клубного дома у метро и парка в расчёте — {m(P0["price"] / 1e3)} тыс. ₽ за м² без отделки, для 1 этажа (магазины) — {m(P0["retail"] / 1e3)} тыс.</p>

<h1>3. Почём земля в Москве</h1>
<img class=full src="file://{os.path.abspath(f"{D}/dev_ha.png")}">
<p class=src>550 млн за 484 м² — {m(ha_offer, 1)} млрд ₽ за гектар, почти как самая дорогая сделка 2026 года (у Павелецкого вокзала, под высокую
застройку). Вместе с общим участком :1863 — {m(550 / 1e3 / (PLOT_ALL / 1e4), 1)} млрд за га, но его делим с соседями. Сделки — «Ведомости» (IBC, Asseto).</p>
<table class=t><thead><tr><th>ID ЦИАН</th><th>Район</th><th class=r>Участок, м²</th><th class=r>Млн ₽</th><th class=r>Разрешено, м²</th>
<th class=r>Тыс. ₽ за м² разреш.</th><th>Что разрешено</th></tr></thead><tbody>{g_rows}</tbody></table>
<p class=src>Нежилые участки и здания под снос в старой Москве на ЦИАН, у которых в объявлении указана разрешённая площадь (ГПЗУ, ГЗК, ТЭП).
Медиана — {m(gmed / 1e3)} тыс. ₽ за м² разрешённой площади. У большинства из них документы уже есть, у нас — нет.</p>
</section>

<section class=pb><h1>4. Сколько девелопер может заплатить за участок</h1>
<img class=full src="file://{os.path.abspath(f"{D}/dev_curve.png")}">
<table class=t style="margin-top:6px"><thead><tr><th>Этажей</th><th class=r>Наземная, м²</th><th class=r>Продаётся, м²</th>
<th class=r>Выручка</th><th class=r>Затраты</th><th class=r>Продажи и прибыль</th><th class=r>Остаток к стройке</th>
<th class=r>Земля сегодня</th><th class=r>С арендой МФЦ 2 года</th></tr></thead><tbody>{rows}</tbody></table>
<p class=src>Суммы — в млн ₽. Выручка от продажи минус стройка (наземная часть {m(P0["cost"] / 1e3)} тыс. ₽/м², технический подземный этаж
{m(P0["ucost"] / 1e3)} тыс.), проектирование и сети 13%, снос, раскопки и плата за смену ВРИ ({m((P0["demo"] + P0["arch"] + P0["vri"]) / 1e6)} млн),
финансирование стройки 9%, продажи 4% и прибыль девелопера 20% выручки. Остаток дисконтирован на {m(P0["years"], 1)} года согласований
под {m(P0["bridge"] * 100)}% годовых. Все допущения можно поменять в таблице Excel.</p>
<h2>Сверка с рынком земли</h2>
<table class=t><tbody>{meth_rows}</tbody></table>
<p class=src>Последняя строка — цена участка с уже готовыми разрешениями. У нас их нет, а получение займёт 2–3 года.</p>
</section>

<section class=pb><h1>5. Что в статье «Ведомостей» важно для нас</h1>
<ul>
<li><b>Спрос на землю есть.</b> За 9 месяцев 2026 г. девелоперы вложили в участки под жильё 211 млрд ₽ (+37%), 85% — в столичном регионе.
В столице, по данным IBC, вложения на уровне прошлого года (95 млрд против 96), сделок меньше (29 против 34).</li>
<li><b>Средняя цена 1 га в Москве — 3,9 млрд ₽</b>, в 2026 году — от 1,4 до 12 млрд за га. Наша цена 550 млн — 11,4 млрд за га, у верхней границы.
Наш участок 484 м² по средней цене стоил бы {m(3.9e9 * PLOT / 1e4 / 1e6)} млн.</li>
<li><b>Больше 70% сделок — участки с утверждёнными параметрами застройки.</b> У нас их нет, плюс охранные зоны. Значит, наш участок для
девелопера стоит дешевле среднего, а не дороже.</li>
<li><b>Земля — 12–13% затрат проекта, с документами — до трети.</b> По этой доле участок под дом на 4–6 этажей стоит около
{m(share(4, 0.125))}–{m(share(6, 0.125))} млн, а с готовыми документами — {m(share(4, 1 / 3))}–{m(share(6, 1 / 3))} млн. Всё ниже 550.</li>
<li><b>Кредит на покупку земли — больше 20% годовых, цикл проекта — 5–8 лет.</b> Каждый год согласований обходится девелоперу дорого,
поэтому за участок без документов он платит с большой скидкой.</li>
<li><b>Продавцы стали сговорчивее, на рынок выходят новые участки.</b> Выбор у девелоперов шире, чем год назад. Предложение 550 млн за
маленький участок без документов на таком фоне — редкая удача для продавца.</li>
</ul>

<h1>6. Вывод</h1>
<ul>
<li><b>Как участок для застройки объект стоит {lo}–{hi} млн.</b> Предложение 550 млн в {m(550 / hi, 1)} раза больше верхней границы,
а оптимистичной оценки ({m(opt6["total"] / 1e6)} млн) — в {m(550 / (opt6["total"] / 1e6), 1)} раза.</li>
<li><b>Здание с МФЦ ({m(CURRENT)} млн) стоит больше, чем земля под ним.</b> Поэтому ориентир для торга — не цена земли, а 383 млн при
текущем использовании. 550 млн выше его на 44%.</li>
<li><b>Почему покупатель может платить 550:</b> он собирает квартал (общий участок :1863 и соседние здания), знает о будущем проекте на
этой территории или ему нужно именно это здание у метро. Если покупатель уже контролирует соседей, наш участок для него ключевой и
цену можно держать.</li>
<li><b>Главные риски:</b> снижение цены после проверки документов («нашли охранную зону», «нет ГПЗУ», «надо платить за ВРИ») или сделка
с условием «если получим разрешения» на 2–3 года. В обоих случаях сделка будет совсем другой.</li>
<li><b>Что сделать сейчас:</b>
  <ul>
  <li>заказать выписки ЕГРН на соседние здания :1735 и :4604 и общий участок :1863 — увидим, кто их владельцы и арендаторы;</li>
  <li>получить ГПЗУ на участок :4 — будем знать, что на нём можно строить, раньше покупателя;</li>
  <li>зафиксировать 550 млн задатком и не соглашаться на условия о разрешениях;</li>
  <li>с МФЦ выбрать вариант на 1 год или 3 года с правом досрочного расторжения — девелоперу нужно свободное здание.</li>
  </ul></li>
</ul>
<p class=src>Источники: «Ведомости», «Девелоперы запасаются землёй» (октябрь 2026), «Цены на землю под строительство жилья в Москве начали падать»
(14.07.2026); «Коммерсантъ» — себестоимость офисного строительства (IBC Real Estate и Stone); объявления ЦИАН на {R.DATE}; выписки ЕГРН.
Плата за смену ВРИ и стоимость раскопок — оценки, точные суммы дают ДГИ и Департамент культурного наследия.</p>
</section>'''
    return f'<!doctype html><html lang="ru"><head><meta charset="utf-8"><style>{css}</style></head><body>{body}</body></html>'


def write_pdf(html):
    open(HTML, 'w').write(html)
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page()
        pg.goto('file://' + os.path.abspath(HTML))
        pg.wait_for_timeout(1200)
        pg.pdf(path=PDF, format='A4', print_background=True, margin=dict(top='17mm', bottom='14mm', left='16mm', right='15mm'),
               display_header_footer=True,
               header_template='<div style="font-family:Carlito,sans-serif;font-size:7.5pt;color:#555;width:100%;padding:0 15mm 0 16mm;'
                               'display:flex;justify-content:space-between"><span>Стромынка, 2: участок для девелопера</span>'
                               f'<span>{R.DATE}</span></div>',
               footer_template='<div style="font-family:Carlito,sans-serif;font-size:8pt;color:#555;width:100%;text-align:center">'
                               '<span class="pageNumber"></span></div>')
        b.close()


if __name__ == '__main__':
    nb, gc = newbuild(), gns_comps()
    curve_png(f'{D}/dev_curve.png')
    ha_png(f'{D}/dev_ha.png')
    nb_png(nb, f'{D}/dev_nb.png')
    rows = write_xlsx(nb, gc)
    write_pdf(build_html(nb, gc))
    for f in FLOORS:
        b = model(f)
        print(f, round(b['res0'] / 1e6), round(b['total'] / 1e6), round(model(f, dict(rent_years=2))['total'] / 1e6))
    print('этажей для 550:', round(floors_for(550e6), 1), 'с МФЦ', round(floors_for(550e6, dict(rent_years=2)), 1),
          'оптимистично', round(floors_for(550e6, OPTIMISTIC), 1), '| опт. 6 эт.', round(model(6, OPTIMISTIC)['total'] / 1e6))
    print('ГНС медиана', round(st.median(r['per'] for r in gc)), 'строки', rows)
