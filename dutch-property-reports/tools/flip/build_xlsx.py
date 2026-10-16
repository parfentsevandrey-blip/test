"""Excel-модель флипа Kernweg 1 с живыми формулами.

python3 -I build_xlsx.py KERNWEG_JSON COMPS_JSON CANDIDATES_JSON OUT_XLSX
"""
import json
import re
import sys
from pathlib import Path

sys.path.append("/root/.local/lib/python3.11/site-packages")
from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import ColorScaleRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402

K = json.loads(Path(sys.argv[1]).read_text())
C = json.loads(Path(sys.argv[2]).read_text())
CANDS = json.loads(Path(sys.argv[3]).read_text())
OUT = Path(sys.argv[4])

NAVY, OCHRE, SOFT = "16233A", "9C7C38", "F5F4F0"
F_HEAD = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
F_TITLE = Font(name="Calibri", bold=True, color=NAVY, size=14)
F_SUB = Font(name="Calibri", italic=True, color="5D6B7C", size=9)
F_INPUT = Font(name="Calibri", color="1F4FB4", size=10)
F_BOLD = Font(name="Calibri", bold=True, size=10)
F_NORM = Font(name="Calibri", size=10)
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_INPUT = PatternFill("solid", fgColor="FFF7E0")
FILL_KEY = PatternFill("solid", fgColor="ECE6DA")
THIN = Side(style="thin", color="DCDDDE")
EUR = '"€" #,##0;[Red]-"€" #,##0'
EUR2 = '"€" #,##0.00'
PCT = "0.0%"
PCT2 = "0.00%"
NUM = "#,##0"

wb = Workbook()


RENAME = {"price": "buy_price", "rate": "loan_rate", "area": "area_vvo"}


def name(ref_sheet, cell, nm):
    nm = RENAME.get(nm, nm)
    wb.defined_names[nm] = DefinedName(nm, attr_text=f"'{ref_sheet}'!${cell[0]}${cell[1:]}"
                                       if cell[1:].isdigit() else f"'{ref_sheet}'!{cell}")


def absref(col, row):
    return f"${col}${row}"


def header(ws, row, values, widths=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="center" if i > 1 else "left", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB


# ---------------------------------------------------------------- допущения
A = wb.active
A.title = "Допущения"
title(A, "Флип-модель: Kernweg 1, Хорн — допущения",
      "Синие значения на жёлтом фоне можно менять: все листы пересчитаются. Суммы без НДС, евро.")
A.column_dimensions["A"].width = 46
A.column_dimensions["B"].width = 16
A.column_dimensions["C"].width = 70
p = K["base_params"]
inputs = [
    ("Покупка", None, None, None),
    ("Цена продавца (k.k.)", K["ask"], EUR, "ask", "funda in business, 16.10.2026"),
    ("Цена покупки — наше предложение", K["price"], EUR, "price", "−23 % к цене продавца; потолок без убытка — см. «Сводка»"),
    ("Площадь, м² VVO (полезная)", K["area"], NUM, "area", "по объявлению: 2.016 м² VVO, 2.155 м² BVO"),
    ("Налог на передачу недвижимости", 0.104, PCT, "ovb", "ст. 14 WBR, нежилая недвижимость, 2026–2027"),
    ("Нотариус и кадастр", 3500, EUR, "notary", "оценка"),
    ("Проверка объекта: техника, грунт, право", 7500, EUR, "dd", "оценка"),
    ("Оценка для банка", 3500, EUR, "valuation", "оценка"),
    ("Консультант по покупке, % цены", 0.01, PCT, "advisory", "оценка"),
    ("Подготовка к сдаче", None, None, None),
    ("Ремонт, € на м²", 20, EUR, "capex_m2", "покраска, свет, мелкий ремонт, разделение на два входа"),
    ("Ремонт, резерв", 2500, EUR, "capex_fixed", ""),
    ("НДС на ремонт (расход, пока аренда без НДС)", 0.21, PCT, "capex_vat", "при аренде с НДС возвращается"),
    ("Месяцев на ремонт", 2, NUM, "capex_months", ""),
    ("Сдача в аренду", None, None, None),
    ("Комиссия брокера, % годовой аренды", 0.15, PCT, "letting_fee", "рынок 10–16 %"),
    ("Реклама и показы", 1500, EUR, "marketing", ""),
    ("Содержание", None, None, None),
    ("Налог OZB, страховка, ремонт: % цены в год", 0.006, PCT2, "holding_pct", ""),
    ("Энергия и охрана в простое, € на м² в год", 3.0, EUR2, "void_m2", ""),
    ("Финансирование", None, None, None),
    ("Кредит, доля цены покупки (LTV)", 0.5, PCT, "ltv", "банк под пустое здание даёт 50–60 %"),
    ("Euribor 3M", K["euribor"], PCT2, "euribor", "euribor-rates.eu, 14.10.2026"),
    ("Маржа банка", K["margin"], PCT2, "bank_margin", "оценка для кредита под пустое здание"),
    ("Комиссия банка, % кредита", 0.01, PCT, "bank_fee", ""),
    ("Продажа", None, None, None),
    ("Брокер на продаже, % цены", 0.0125, PCT2, "sale_fee", ""),
    ("Юрист и реклама продажи", 2500, EUR, "sale_fixed", ""),
    ("Доля экономии покупателя по ст. 13 WBR, которую забираем в цене", 0.5, PCT, "art13_share",
     "только при продаже до 6 месяцев после нашей покупки"),
    ("Налог на прибыль (VPB) 2026", None, None, None),
    ("Ставка до порога", 0.19, PCT, "vpb_low", "KVK, 2026"),
    ("Ставка выше порога", 0.258, PCT2, "vpb_high", ""),
    ("Порог", 200000, EUR, "vpb_step", ""),
]
row = 4
for item in inputs:
    label, value, fmt, *rest = item
    if value is None:
        A.cell(row=row, column=1, value=label).font = Font(bold=True, color=OCHRE, size=10)
        row += 1
        continue
    nm, note = rest
    A.cell(row=row, column=1, value=label).font = F_NORM
    c = A.cell(row=row, column=2, value=value)
    c.font, c.fill, c.number_format = F_INPUT, FILL_INPUT, fmt
    A.cell(row=row, column=3, value=note).font = F_SUB
    name("Допущения", f"B{row}", nm)
    row += 1
A.cell(row=row, column=1, value="Ставка кредита").font = F_BOLD
A.cell(row=row, column=2, value="=euribor+bank_margin").number_format = PCT2
name("Допущения", f"B{row}", "rate")
row += 2

# сценарии
A.cell(row=row, column=1, value="Сценарии").font = Font(bold=True, color=OCHRE, size=10)
row += 1
scen_names = ["Пессимистичный", "Базовый", "Оптимистичный", "Ускоренный, 6 мес."]
for i, v in enumerate(["Параметр"] + scen_names, start=1):
    c = A.cell(row=row, column=i if i == 1 else i + 2, value=v)
    c.font, c.fill = F_HEAD, FILL_HEAD
for col in "DEFG":
    A.column_dimensions[col].width = 15
row += 1
scen_rows = [
    ("Аренда, множитель к оценке", [0.9, 1.0, 1.1, 1.0], "0.00", "rent_mult"),
    ("Доходность выхода (BAR k.k.)", [0.09, 0.0825, 0.0775, 0.0825], PCT2, "exit_bar"),
    ("Месяц начала аренды", [8, 5, 4, 3], NUM, "lease_month"),
    ("Арендные каникулы, мес.", [4, 3, 2, 2], NUM, "rent_free"),
    ("Месяц продажи (передача у нотариуса)", [15, 12, 10, 6], NUM, "sale_month"),
]
A.cell(row=row - 1, column=2, value="Активный").font = F_HEAD
A.cell(row=row - 1, column=2).fill = FILL_HEAD
sel_row = row + len(scen_rows)
for label, vals, fmt, nm in scen_rows:
    A.cell(row=row, column=1, value=label).font = F_NORM
    for j, v in enumerate(vals):
        c = A.cell(row=row, column=4 + j, value=v)
        c.font, c.fill, c.number_format = F_INPUT, FILL_INPUT, fmt
    c = A.cell(row=row, column=2, value=f"=INDEX(D{row}:G{row},scenario)")
    c.number_format, c.font = fmt, F_BOLD
    name("Допущения", f"B{row}", nm)
    row += 1
A.cell(row=row, column=1, value="Активный сценарий: 1 — пессимистичный, 2 — базовый, 3 — оптимистичный, 4 — ускоренный").font = F_BOLD
c = A.cell(row=row, column=2, value=2)
c.font, c.fill = F_INPUT, FILL_INPUT
name("Допущения", f"B{row}", "scenario")

# ---------------------------------------------------------------- аренда
R = wb.create_sheet("Аренда")
title(R, "Оценка рыночной аренды по частям здания",
      "Ставки — по предложениям в Хорне, Зваге, Энкхёйзене и Пюрмеренде на funda in business, 16.10.2026")
header(R, 4, ["Часть здания", "м² VVO", "€ за м² в год", "Аренда в год", "Основание ставки"], [40, 10, 14, 16, 70])
basis = {
    "Шоурум, первый этаж": "витрина у въезда в Хорн 80, разрешена торговля; мелкие помещения в промзоне — € 85–100",
    "Шоурум и витрины, антресоль": "второй уровень шоурума, без отдельного входа",
    "Офисы, два этажа": "офисы в деловых районах Хорна — € 120–150",
    "Кантина и подсобные": "вспомогательные помещения",
    "Мастерские и склады, первый этаж": "большие залы в Хорне и окрестностях — € 57–75",
    "Склады и мастерские, второй этаж": "склад на этаже без ворот",
}
r0 = 5
for i, (label, area, rate) in enumerate(K["parts"]):
    rr = r0 + i
    R.cell(row=rr, column=1, value=label)
    R.cell(row=rr, column=2, value=area).number_format = NUM
    c = R.cell(row=rr, column=3, value=rate)
    c.font, c.fill, c.number_format = F_INPUT, FILL_INPUT, EUR
    R.cell(row=rr, column=4, value=f"=B{rr}*C{rr}").number_format = EUR
    R.cell(row=rr, column=5, value=basis.get(label, "")).font = F_SUB
rt = r0 + len(K["parts"])
R.cell(row=rt, column=1, value="Итого").font = F_BOLD
R.cell(row=rt, column=2, value=f"=SUM(B{r0}:B{rt - 1})").number_format = NUM
R.cell(row=rt, column=3, value=f"=D{rt}/B{rt}").number_format = EUR2
c = R.cell(row=rt, column=4, value=f"=SUM(D{r0}:D{rt - 1})")
c.number_format, c.font = EUR, F_BOLD
name("Аренда", f"D{rt}", "rent_base")
R.cell(row=rt + 2, column=1, value="Аренда с учётом сценария").font = F_BOLD
R.cell(row=rt + 2, column=4, value="=rent_base*rent_mult").number_format = EUR
name("Аренда", f"D{rt + 2}", "rent_eff")
R.cell(row=rt + 3, column=1, value="К цене продавца")
R.cell(row=rt + 3, column=4, value="=rent_eff/ask").number_format = PCT2
R.cell(row=rt + 4, column=1, value="К цене покупки")
R.cell(row=rt + 4, column=4, value="=rent_eff/price").number_format = PCT2

# ---------------------------------------------------------------- сводка (расчётный блок)
S = wb.create_sheet("Сводка", 0)
title(S, "Флип-модель: Kernweg 1, Хорн — сводка",
      "Купить пустым → сдать → продать сданным инвестору. Все цифры — формулы от листов «Допущения» и «Аренда».")
S.column_dimensions["A"].width = 52
S.column_dimensions["B"].width = 18
S.column_dimensions["C"].width = 60
calc = [
    ("Покупка", None),
    ("Цена покупки", "=price", EUR, "acq_price"),
    ("Налог на передачу 10,4 %", "=price*ovb", EUR, "acq_ovb"),
    ("Нотариус и кадастр", "=notary", EUR, None),
    ("Проверка объекта", "=dd", EUR, None),
    ("Оценка для банка", "=IF(ltv>0,valuation,0)", EUR, None),
    ("Консультант по покупке", "=price*advisory", EUR, None),
    ("Комиссия банка", "=loan*bank_fee", EUR, None),
    ("Итого на покупку", "=SUM(B{a}:B{b})", EUR, "acq_total"),
    ("Кредит", "=price*ltv", EUR, "loan"),
    ("Ремонт с НДС", "=(capex_m2*area+capex_fixed)*(1+capex_vat)", EUR, "capex_total"),
    ("Продажа", None),
    ("Аренда в год (сценарий)", "=rent_eff", EUR, None),
    ("Первый оплачиваемый месяц", "=lease_month+rent_free", NUM, "first_paid"),
    ("Стоимость сданного здания: аренда / доходность для инвестора", "=rent_eff/exit_bar", EUR, "value_let"),
    ("Минус не закончившиеся каникулы", "=MAX(0,first_paid-1-sale_month)*rent_eff/12", EUR, "free_left"),
    ("Плюс доля экономии покупателя по ст. 13 (продажа до 6 мес.)",
     "=IF(sale_month<=6,art13_share*ovb*price/(1+ovb),0)", EUR, "uplift"),
    ("Цена продажи", "=value_let-free_left+uplift", EUR, "sale_price"),
    ("Расходы на продажу", "=sale_price*sale_fee+sale_fixed", EUR, "sale_cost"),
    ("Результат", None),
    ("Доходы: аренда за период и продажа", "=SUM('Денежный поток'!B10:Z10)+SUM('Денежный поток'!B13:Z13)", EUR, "income"),
    ("Все затраты: цена, налоги, ремонт, проценты, комиссии", "=-(SUM('Денежный поток'!B5:Z9)+SUM('Денежный поток'!B11:Z11)+SUM('Денежный поток'!B14:Z14))",
     EUR, "cost_total"),
    ("Прибыль до налога", "=SUM('Денежный поток'!B15:Z15)", EUR, "profit"),
    ("Заработок на каждые € 100 затрат", "=profit/cost_total*100", EUR2, None),
    ("Налог на прибыль VPB", "=IF(profit<=0,0,MIN(profit,vpb_step)*vpb_low+MAX(0,profit-vpb_step)*vpb_high)", EUR, "vpb"),
    ("Прибыль после налога", "=profit-vpb", EUR, None),
    ("Своих денег нужно сверх кредита (максимум)", "=-MIN('Денежный поток'!B16:Z16)", EUR, "equity"),
    ("Прибыль на свои деньги до налога", "=profit/equity", PCT, None),
    ("Прибыль на свои деньги после налога", "=(profit-vpb)/equity", PCT, None),
    ("Доходность своих денег в пересчёте на год (IRR), до налога", "=(1+IRR('Денежный поток'!B15:Z15,0.01))^12-1", PCT, None),
    ("Доходность своих денег в пересчёте на год (IRR), после налога", "=(1+IRR('Денежный поток'!B18:Z18,0.01))^12-1", PCT, None),
    ("Сколько вернётся на € 1 своих денег, после налога", "=(equity+profit-vpb)/equity", EUR2, None),
    ("Глазами покупателя", None),
    ("Налог на передачу, который платит покупатель",
     "=IF(sale_month<=6,ovb*(sale_price-price),ovb*sale_price)", EUR, "buyer_ovb"),
    ("Его полная цена", "=sale_price+buyer_ovb", EUR, "buyer_total"),
    ("Его доходность на все расходы", "=rent_eff/buyer_total", PCT2, None),
]
notes = {
    "Цена продажи": "продажа сданного объекта инвестору; расходы на сделку платит покупатель (k.k.)",
    "Своих денег нужно сверх кредита (максимум)": "сколько вложить из своего кармана в самый затратный момент проекта",
    "Его доходность на все расходы": "сравнить с 8,5–8,9 % у сданных объектов в Хорне",
    "Заработок на каждые € 100 затрат": "прибыль до налога, делённая на все затраты",
    "Доходность своих денег в пересчёте на год (IRR), до налога": "учитывает, сколько месяцев деньги заняты в проекте",
    "Сколько вернётся на € 1 своих денег, после налога": "вложенное плюс чистая прибыль, делённые на вложенное",
    "Плюс доля экономии покупателя по ст. 13 (продажа до 6 мес.)":
        "ст. 13 WBR: покупатель платит 10,4 % только с прироста цены",
}
row = 4
acq_start = None
for item in calc:
    label, formula = item[0], item[1]
    if formula is None:
        S.cell(row=row, column=1, value=label).font = Font(bold=True, color=OCHRE, size=11)
        row += 1
        continue
    fmt, nm = item[2], item[3]
    if label == "Цена покупки":
        acq_start = row
    if "{a}" in formula:
        formula = formula.format(a=acq_start, b=row - 1)
    S.cell(row=row, column=1, value=label).font = F_NORM
    c = S.cell(row=row, column=2, value=formula)
    c.number_format = fmt
    if label in ("Прибыль до налога", "Прибыль после налога", "Цена продажи", "Итого на покупку",
                 "Своих денег нужно сверх кредита (максимум)",
                 "Доходность своих денег в пересчёте на год (IRR), до налога"):
        c.font = F_BOLD
        c.fill = FILL_KEY
        S.cell(row=row, column=1).font = F_BOLD
    if label in notes:
        S.cell(row=row, column=3, value=notes[label]).font = F_SUB
    if nm:
        name("Сводка", f"B{row}", nm)
    row += 1

row += 1
S.cell(row=row, column=1, value="Потолок цены покупки (подбор моделью, базовые допущения)").font = Font(bold=True, color=OCHRE, size=11)
row += 1
for label, key, sub in [("Прибыль 0 — дороже не покупать, продажа через 12 мес.", "breakeven", "12"),
                        ("€ 10 на каждые € 100 затрат, продажа через 12 мес.", "m10", "12"),
                        ("€ 15 на каждые € 100 затрат, продажа через 12 мес.", "m15", "12"),
                        ("Прибыль 0, продажа через 6 мес. (ст. 13)", "breakeven", "6"),
                        ("€ 10 на каждые € 100 затрат, продажа через 6 мес.", "m10", "6")]:
    S.cell(row=row, column=1, value=label)
    S.cell(row=row, column=2, value=round(K["bids"][sub][key])).number_format = EUR
    S.cell(row=row, column=3, value=f"={get_column_letter(2)}{row}/ask-1").number_format = PCT
    row += 1
S.cell(row=row, column=1, value="Доходность для инвестора, при которой прибыль = 0 (12 мес.)")
S.cell(row=row, column=2, value=K["bids"]["12"]["bar_be"]).number_format = PCT2
row += 1
S.cell(row=row, column=1, value="Аренда, при которой прибыль = 0, € за м² (12 мес.)")
S.cell(row=row, column=2, value=K["bids"]["12"]["rent_be"]).number_format = EUR2
row += 1
S.cell(row=row, column=1, value="Значения подобраны Python-моделью; в Excel их можно перепроверить подбором параметра.").font = F_SUB

# ---------------------------------------------------------------- денежный поток
D = wb.create_sheet("Денежный поток", 2)
title(D, "Денежный поток по месяцам", "Месяц 0 — передача у нотариуса и оплата; продажа — в конце месяца продажи")
D.column_dimensions["A"].width = 38
labels = ["Месяц", "В работе (до продажи)", "Покупка и расходы на сделку", "Ремонт", "Содержание",
          "Энергия и охрана в простое", "Сдача: брокер и реклама", "Аренда", "Проценты", "Кредит: получен / погашен",
          "Продажа", "Расходы на продажу", "Поток до налога", "Накопленный поток", "Налог на прибыль",
          "Поток после налога"]
for i, label in enumerate(labels):
    c = D.cell(row=3 + i, column=1, value=label)
    c.font = F_BOLD if label in ("Месяц", "Поток до налога", "Накопленный поток", "Поток после налога") else F_NORM
for m in range(25):
    col = get_column_letter(2 + m)
    D.column_dimensions[col].width = 11
    mc = f"{col}$3"
    D[f"{col}3"] = m
    D[f"{col}3"].font, D[f"{col}3"].fill = F_HEAD, FILL_HEAD
    on = f"({mc}<=sale_month)"
    D[f"{col}4"] = f"=IF({mc}<=sale_month,1,0)"
    D[f"{col}5"] = f"=IF({mc}=0,-acq_total,0)"
    D[f"{col}6"] = f"=IF(AND({mc}>=1,{mc}<=capex_months,{mc}<=sale_month),-capex_total/capex_months,0)"
    D[f"{col}7"] = f"=IF(AND({mc}>=1,{mc}<=sale_month),-price*holding_pct/12,0)"
    D[f"{col}8"] = f"=IF(AND({mc}>=1,{mc}<lease_month,{mc}<=sale_month),-void_m2*area/12,0)"
    D[f"{col}9"] = f"=IF(AND({mc}=lease_month,{mc}<=sale_month),-(rent_eff*letting_fee+marketing),0)"
    D[f"{col}10"] = f"=IF(AND({mc}>=1,{mc}>=first_paid,{mc}<=sale_month),rent_eff/12,0)"
    D[f"{col}11"] = f"=IF(AND({mc}>=1,{mc}<=sale_month),-loan*rate/12,0)"
    D[f"{col}12"] = f"=IF({mc}=0,loan,0)-IF({mc}=sale_month,loan,0)"
    D[f"{col}13"] = f"=IF({mc}=sale_month,sale_price,0)"
    D[f"{col}14"] = f"=IF({mc}=sale_month,-sale_cost,0)"
    D[f"{col}15"] = f"=SUM({col}5:{col}14)"
    prev = get_column_letter(1 + m)
    D[f"{col}16"] = f"={col}15" if m == 0 else f"={prev}16+{col}15"
    D[f"{col}17"] = f"=IF({mc}=sale_month,-vpb,0)"
    D[f"{col}18"] = f"={col}15+{col}17"
    for r in range(4, 19):
        D[f"{col}{r}"].number_format = NUM if r == 4 else EUR
    for r in (15, 16, 18):
        D[f"{col}{r}"].font = F_BOLD
D.freeze_panes = "B4"

# ---------------------------------------------------------------- сценарии
SC = wb.create_sheet("Сценарии")
title(SC, "Сценарии и лестница цены покупки", "Рассчитано той же моделью; активный сценарий переключается на листе «Допущения»")
names = list(K["scenarios"])
header(SC, 4, ["Показатель"] + names, [44, 18, 18, 18, 18])
rows = [
    ("Аренда в год", lambda r: r["rent"], EUR),
    ("Доходность для инвестора", lambda r: r["p"]["exit_bar"], PCT2),
    ("Месяц начала аренды", lambda r: r["p"]["lease_month"], NUM),
    ("Месяц продажи", lambda r: r["p"]["sale_month"], NUM),
    ("Цена продажи", lambda r: r["sale"], EUR),
    ("Все затраты", lambda r: r["cost_total"], EUR),
    ("Прибыль до налога", lambda r: r["profit"], EUR),
    ("Заработок на € 100 затрат", lambda r: r["margin"] * 100, EUR2),
    ("Прибыль после налога", lambda r: r["profit_after"], EUR),
    ("Своих денег нужно сверх кредита", lambda r: r["equity"], EUR),
    ("Прибыль на свои деньги после налога", lambda r: r["roe_after"], PCT),
    ("В пересчёте на год до налога (IRR)", lambda r: r["irr"], PCT),
    ("В пересчёте на год после налога (IRR)", lambda r: r["irr_after"], PCT),
]
for i, (label, fn, fmt) in enumerate(rows):
    SC.cell(row=5 + i, column=1, value=label)
    for j, nm in enumerate(names):
        c = SC.cell(row=5 + i, column=2 + j, value=fn(K["scenarios"][nm]))
        c.number_format = fmt
row = 5 + len(rows) + 2
SC.cell(row=row, column=1, value="Лестница цены покупки (базовые допущения)").font = Font(bold=True, color=OCHRE, size=11)
row += 1
for i, v in enumerate(["Цена покупки", "Скидка к цене продавца", "€ за м² VVO", "Прибыль, 12 мес.", "На € 100 затрат, 12 мес.",
                       "В год, 12 мес.", "Прибыль, 6 мес.", "На € 100 затрат, 6 мес.", "В год, 6 мес."], start=1):
    c = SC.cell(row=row, column=i, value=v)
    c.font, c.fill = F_HEAD, FILL_HEAD
    SC.column_dimensions[get_column_letter(i)].width = max(SC.column_dimensions[get_column_letter(i)].width or 0, 15)
for l in K["ladder"]:
    row += 1
    for i, (v, fmt) in enumerate([(l["price"], EUR), (l["disc"], PCT), (l["pm2"], EUR), (l["p12"], EUR), (l["m12"] * 100, EUR2),
                                  (l["irr12"], PCT), (l["p6"], EUR), (l["m6"] * 100, EUR2), (l["irr6"], PCT)], start=1):
        c = SC.cell(row=row, column=i, value=v)
        c.number_format = fmt
        if l["price"] == K["price"]:
            c.fill = FILL_KEY

# ---------------------------------------------------------------- чувствительность
SE = wb.create_sheet("Чувствительность")
title(SE, "Чувствительность прибыли до налога", "Продажа через 12 месяцев; остальные допущения базовые")
SE.column_dimensions["A"].width = 30


def grid(top, caption, row_label, row_vals, row_fmt, cols, col_fmt, cells):
    SE.cell(row=top, column=1, value=caption).font = Font(bold=True, color=OCHRE, size=11)
    SE.cell(row=top + 1, column=1, value=row_label).font = F_HEAD
    SE.cell(row=top + 1, column=1).fill = FILL_HEAD
    for j, v in enumerate(cols):
        c = SE.cell(row=top + 1, column=2 + j, value=v)
        c.font, c.fill, c.number_format = F_HEAD, FILL_HEAD, col_fmt
        SE.column_dimensions[get_column_letter(2 + j)].width = 13
    for i, rv in enumerate(row_vals):
        c = SE.cell(row=top + 2 + i, column=1, value=rv)
        c.number_format, c.font = row_fmt, F_BOLD
        for j, v in enumerate(cells[i]):
            SE.cell(row=top + 2 + i, column=2 + j, value=v).number_format = EUR
    rng = f"B{top + 2}:{get_column_letter(1 + len(cols))}{top + 1 + len(row_vals)}"
    SE.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="E6B8AF", mid_type="num",
                                                      mid_value=0, mid_color="FFFFFF", end_type="max",
                                                      end_color="B6D7A8"))
    return top + 3 + len(row_vals)


g = K["grid_price_bar"]
top = grid(4, "Прибыль: цена покупки (строки) и доходность, которую требует инвестор (столбцы)", "Цена \\ доходность", g["prices"], EUR, g["bars"], PCT2,
           g["profit"])
g = K["grid_rent_bar"]
top = grid(top + 1, "Прибыль: аренда выше или ниже оценки (строки) и доходность для инвестора (столбцы)", "Аренда \\ доходность", g["rents"], "0%",
           g["bars"], PCT2, g["profit"])
g = K["grid_lease_sale"]
cells = [[(c["profit"] if c else None) for c in row] for row in g["cells"]]
top = grid(top + 1, "Месяц начала аренды (строки) × месяц продажи (столбцы)", "Аренда \\ продажа", g["leases"], NUM,
           g["sales"], NUM, cells)
SE.cell(row=top + 1, column=1, value="Доля экономии покупателя по ст. 13, которую забираем (продажа через 6 мес.)").font = Font(bold=True, color=OCHRE, size=11)
for i, (share, prof) in enumerate(K["art13_share"]):
    SE.cell(row=top + 2, column=2 + i, value=share).number_format = "0%"
    SE.cell(row=top + 3, column=2 + i, value=prof).number_format = EUR
SE.cell(row=top + 2, column=1, value="Доля")
SE.cell(row=top + 3, column=1, value="Прибыль до налога")

# ---------------------------------------------------------------- план Б
B = wb.create_sheet("План Б")
title(B, "План Б: если быстро продать не получается", "Формулы от листов «Допущения», «Аренда» и «Сводка»")
B.column_dimensions["A"].width = 52
B.column_dimensions["B"].width = 18
B.column_dimensions["C"].width = 58
plan = [
    ("Держать сданным и рефинансировать", None),
    ("Расходы владельца, % аренды", 0.125, PCT, "hold_opex", "налог OZB владельца, страховка, ремонт, управление"),
    ("Новый кредит, % стоимости со сдачей", 0.55, PCT, "refi_ltv", ""),
    ("Чистый доход в год", "=rent_eff*(1-hold_opex)", EUR, "hold_noi", ""),
    ("Вложено всего", "=acq_total+capex_total+rent_eff*letting_fee+marketing", EUR, "hold_invested", ""),
    ("Чистая доходность на вложенное", "=hold_noi/hold_invested", PCT2, None, ""),
    ("Новый кредит", "=refi_ltv*value_let", EUR, "hold_loan", ""),
    ("Проценты в год", "=hold_loan*rate", EUR, "hold_int", ""),
    ("Капитал после рефинансирования", "=hold_invested-hold_loan", EUR, "hold_eq", ""),
    ("Денежный поток в год", "=hold_noi-hold_int", EUR, "hold_cash", ""),
    ("Денежная доходность капитала", "=hold_cash/hold_eq", PCT, None, ""),
    ("Разделить на две части и продать пользователям", None),
    ("Цена продажи, € за м² VVO", 850, EUR, "split_pm2", "Хорн 80: Electronplein 1 — € 922–999, De Marowijne 45 в Зваге — € 1.041"),
    ("Разделение: нотариус, кадастр, счётчики, противопожарная стена, второй вход", 90000, EUR, "split_cost", "оценка"),
    ("Срок, месяцев", 10, NUM, "split_months", ""),
    ("Брокер на продаже, %", 0.02, PCT, "split_fee", ""),
    ("Выручка", "=split_pm2*area", EUR, "split_gross", ""),
    ("Расходы на продажу двух частей", "=split_gross*split_fee+2*5000", EUR, "split_sale_cost", ""),
    ("Проценты за срок", "=loan*rate*split_months/12", EUR, "split_int", ""),
    ("Содержание и простой за срок", "=(price*holding_pct+void_m2*area)*split_months/12", EUR, "split_hold", ""),
    ("Все затраты", "=acq_total+capex_total+split_cost+split_int+split_hold+split_sale_cost", EUR, "split_total", ""),
    ("Прибыль до налога", "=split_gross-split_total", EUR, None, ""),
    ("Маржа на затраты", "=(split_gross-split_total)/split_total", PCT, None, ""),
]
row = 4
for item in plan:
    if item[1] is None:
        B.cell(row=row, column=1, value=item[0]).font = Font(bold=True, color=OCHRE, size=11)
        row += 1
        continue
    label, value, fmt, nm, note = item
    B.cell(row=row, column=1, value=label)
    c = B.cell(row=row, column=2, value=value)
    c.number_format = fmt
    if not str(value).startswith("="):
        c.font, c.fill = F_INPUT, FILL_INPUT
    B.cell(row=row, column=3, value=note).font = F_SUB
    if nm:
        name("План Б", f"B{row}", nm)
    row += 1

# ---------------------------------------------------------------- конкуренты
CO = wb.create_sheet("Конкуренты")
title(CO, "Сопоставимые объекты: Хорн и Западная Фрисландия", "funda in business, выгрузка через Scrapling 16.10.2026")
row = 4


def table(row, caption, cols, widths, data):
    CO.cell(row=row, column=1, value=caption).font = Font(bold=True, color=OCHRE, size=11)
    row += 1
    for i, v in enumerate(cols, start=1):
        c = CO.cell(row=row, column=i, value=v)
        c.font, c.fill = F_HEAD, FILL_HEAD
        CO.column_dimensions[get_column_letter(i)].width = max(CO.column_dimensions[get_column_letter(i)].width or 0, widths[i - 1])
    for rec in data:
        row += 1
        for i, (v, fmt) in enumerate(rec, start=1):
            c = CO.cell(row=row, column=i, value=v)
            if fmt == "link":
                c.hyperlink = v
                c.font = Font(color="1F4FB4", underline="single", size=9)
            elif fmt:
                c.number_format = fmt
    return row + 2


sale = [[(c["addr"], None), (c["town"], None), (c["area"], NUM), (c["price"], EUR), (c["pm2"], EUR),
         (c["year"], None), (c["status"], None), (c["since"], None), (c["url"], "link")] for c in C["sale"]
        if c["pm2"] >= 300 and "aanbouw" not in str(c["year"]) and c["addr"] != "Westrak 11"]   # Westrak 11 — участок под стройку
row = table(row, "Продажа пустых зданий от 400 м²", ["Адрес", "Город", "м²", "Цена", "€ за м²", "Год", "Статус",
                                                      "На продаже", "Ссылка"], [30, 14, 9, 14, 10, 12, 18, 12, 60], sale)
rent = [[(c["addr"], None), (c["town"], None), (c["area"], NUM), (c["rent_m2"], EUR2), (c["area"] * c["rent_m2"], EUR),
         (c["year"], None), (None, None), (c["since"], None), (c["url"], "link")] for c in C["rent"] if c["area"] >= 300]
row = table(row, "Аренда залов от 300 м²", ["Адрес", "Город", "м²", "€ за м² в год", "В год", "Год", "", "На рынке",
                                            "Ссылка"], [30, 14, 9, 14, 10, 12, 18, 12, 60], rent)
inv = [[(c["addr"], None), (c["town"], None), (c["area"], NUM), (c["price"], EUR), (c["rent"], EUR),
        (c["giy"] / 100, PCT2), (c["year"], None), (c["since"], None), (c["url"], "link")] for c in C["inv"]]
inv.append([("Neutronweg 3", None), ("Хорн", None), (2117, NUM), (1650000, EUR), (141000, EUR), (0.0855, PCT2),
            ("около 1978", None), ("снят с продажи", None),
            ("https://www.fundainbusiness.nl/belegging/hoorn-nh/object-43660627-neutronweg-3/", "link")])
row = table(row, "Сданные объекты на продаже — доходность покупателя (k.k.)",
            ["Адрес", "Город", "м²", "Цена", "Аренда в год", "Доходность", "Год", "На продаже", "Ссылка"],
            [30, 14, 9, 14, 10, 12, 18, 12, 60], inv)
cand = [[(c["name"], None), (c["price"], EUR), (c["rent"], EUR), (c["bar"], PCT2), (c["profit_ask"], EUR),
         (-c["disc_be"], PCT), (-c["disc_m10"], PCT), (c["price_m10"], EUR)] for c in CANDS]
row = table(row, "Другие здания по той же схеме: купить по цене продавца, сдать, продать через 12 месяцев",
            ["Объект", "Цена продавца", "Аренда в год, оценка", "Доходность для инвестора",
             "Итог по цене продавца", "Сбить цену, чтобы выйти в ноль", "Сбить цену ради € 10 на € 100 затрат",
             "Цена для € 10 на € 100 затрат"], [30, 14, 9, 14, 10, 12, 18, 12], cand)

# ---------------------------------------------------------------- источники
SR = wb.create_sheet("Источники")
title(SR, "Источники и правила")
SR.column_dimensions["A"].width = 40
SR.column_dimensions["B"].width = 110
src = [
    ("Объект", "https://www.fundainbusiness.nl/bedrijfshal/hoorn-nh/object-43955233-kernweg-1/ — € 1.495.000 k.k., на продаже больше 6 месяцев"),
    ("Налог на передачу", "Wet op belastingen van rechtsverkeer, ст. 14: 10,4 % для нежилой недвижимости (редакции 2026-01-01 и 2027-01-01)"),
    ("Перепродажа до 6 месяцев", "Там же, ст. 13: при покупке в течение 6 месяцев после предыдущей покупки база уменьшается на сумму, с которой налог уже уплачен"),
    ("Налог на прибыль", "VPB 2026: 19 % до € 200.000, 25,8 % сверх (KVK)"),
    ("Ставка", "Euribor 3M 2,628 % на 14.10.2026 (euribor-rates.eu)"),
    ("Рынок", "NVM, бизнес-недвижимость, 2-й квартал 2026: спрос в Северной Голландии за полугодие на уровне прошлого года, предложение 5 % фонда"),
    ("Рынок", "Cushman & Wakefield, 2-й квартал 2026: на второстепенных локациях — пустующие площади и уступки арендаторам"),
    ("Рынок", "JLL, лёгкая промышленность, 2-й квартал 2026: спрос за полугодие 887.200 м², в основном небольшие помещения, пустует 2,2 %"),
    ("Сопоставимые объекты", "funda in business: Хорн, Зваг, Энкхёйзен, Вогнюм, Медемблик, Пюрмеренд, Херхюговард, Алкмар; выгрузка Scrapling 16.10.2026"),
]
for i, (a, b) in enumerate(src, start=4):
    SR.cell(row=i, column=1, value=a).font = F_BOLD
    SR.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True)

TOKEN = re.compile(r"(?<![A-Za-z_'!$])\b(price|rate|area)\b(?![(_A-Za-z])")
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    for line in ws.iter_rows():
        for cell in line:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = TOKEN.sub(lambda m: RENAME[m.group(1)], cell.value)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
