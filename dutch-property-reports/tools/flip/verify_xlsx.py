"""Сверка пересчитанной Excel-модели с выходом flipobj: python3 -I verify_xlsx.py XLSX MODEL_JSON [SCEN]"""
import json
import sys
sys.path.append("/root/.local/lib/python3.11/site-packages")
from openpyxl import load_workbook

wb = load_workbook(sys.argv[1], data_only=True)
M = json.loads(open(sys.argv[2]).read())
scen = sys.argv[3] if len(sys.argv) > 3 else "Базовый"
r = M["base"] if scen == "Базовый" else M["scenarios"][scen]
ws = wb["Сводка"]
vals = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).value for i in range(1, ws.max_row + 1)}
checks = [("Цена продажи", r["sale"]), ("Прибыль до налога", r["profit"]), ("Прибыль после налога", r["profit_after"]),
          ("Вложено своих денег на пике", r["equity"]), ("В пересчёте на год (IRR), после налога", r["irr_after"]),
          ("Налог на передачу, который платит инвестор", r["buyer_ovb"])]
bad = 0
for k, v in checks:
    x = vals.get(k)
    ok = x is not None and abs(x - v) <= max(1.0, abs(v) * 1e-4)
    bad += not ok
    print(f"{'OK ' if ok else 'ERR'} {k}: excel {x} model {v}")
cost = [v for k, v in vals.items() if k and str(k).startswith("Все затраты")][0]
print(("OK " if abs(cost - r["cost_total"]) < 1 else "ERR"), "Все затраты", cost, r["cost_total"])
sys.exit(1 if bad else 0)
