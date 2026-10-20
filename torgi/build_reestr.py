"""Build the tender registry (reestr_torgi.xlsx) in the agreed final format.

Usage: python3 torgi/build_reestr.py [torgi/lots.json] [reestr_torgi.xlsx]

The format mirrors the reference file approved on 20.10.2026: one sheet «Реестр», an Excel table
with sort/filter arrows, 19 columns, coral theme. Formulas («Цена за м²», «Дней до конца приёма»)
are recalculated by Excel on open.
"""
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "lots.json"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE.parent / "reestr_torgi.xlsx"

FONT = "Arial"
INK, GREY = "1F2328", "9AA0A6"
HEAD, ACCENT, ZEBRA, BAR = "2B2D31", "E0533D", "FAFAF9", "F4A493"
RULE = Side(style="thin", color="E6E7EA")
SEP = Side(style="medium", color=ACCENT)

# key, header, width, kind (text/ctext wrap; ctext has a centred header)
COLS = [
    ("n", "№", 5, "num"),
    ("type", "Тип объекта", 20, "text"),
    ("address", "Адрес", 40, "text"),
    ("ao", "Округ", 10.6640625, "ctext"),
    ("district", "Район", 17, "text"),
    ("metro", "Метро", 17, "text"),
    ("area", "Площадь, м²", 14, "area"),
    ("land", "Участок, м²", 14, "area0"),
    ("floor", "Этаж", 12, "text"),
    ("floors", "Этажность", 15, "num"),
    ("okn", "ОКН", 8.77734375, "ctext"),
    ("price", "Начальная цена, ₽", 17, "money"),
    ("ppm", "Цена за м², ₽", 15, "money"),
    ("deposit", "Задаток, ₽", 15, "money2"),
    ("step", "Шаг аукциона, ₽", 14, "money2"),
    ("app_to", "Приём заявок до", 19.33203125, "date"),
    ("days", "Дней до конца приёма", 11, "num"),
    ("auction", "Дата торгов", 19.44140625, "date"),
    ("url", "Карточка", 11, "link"),
]
BLOCK_START = {"area", "price", "url"}  # coral vertical separators between column blocks
C = {k: i for i, (k, *_) in enumerate(COLS, 1)}
DASH = {"type", "address", "ao", "district", "metro", "floor", "okn", "app_to", "auction"}
FORMATS = {"money": "#,##0", "money2": "#,##0.00", "area": "#,##0.0", "area0": "#,##0",
           "date": "dd.mm.yyyy hh:mm"}


def font(bold=False, color=INK, size=10, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, underline=underline)


def fill(color):
    return PatternFill("solid", fgColor=color)


def parse_dt(v):
    m = re.search(r"(\d{2})\.(\d{2})\.(\d{4})(?:\s+(\d{2}):(\d{2}))?", v or "")
    if not m:
        return None
    d, mo, y, h, mi = m.groups()
    return datetime(int(y), int(mo), int(d), int(h or 0), int(mi or 0))


def head_align(kind):
    if kind in ("money", "money2", "area", "area0"):
        return "right"
    return "left" if kind == "text" else "center"


lots = json.loads(SRC.read_text(encoding="utf-8"))["lots"]
lots.sort(key=lambda r: r.get("price") or 0)

wb = Workbook()
ws = wb.active
ws.title = "Реестр"
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = ACCENT
first, last = 2, 1 + len(lots)

ws.row_dimensions[1].height = 46.05
for k, title, width, kind in COLS:
    ws.column_dimensions[L(C[k])].width = width
    c = ws.cell(1, C[k], title)
    c.font = font(bold=True, color="FFFFFF")
    c.fill = fill(HEAD)
    c.border = Border(bottom=Side(style="thick", color=ACCENT), left=SEP if k in BLOCK_START else None)
    c.alignment = Alignment(horizontal=head_align(kind), vertical="center", wrap_text=True, indent=1)

for n, lot in enumerate(lots, 1):
    row = n + 1
    ws.row_dimensions[row].height = 40.05
    values = dict(lot, n=n, app_to=parse_dt(lot.get("app_to")), auction=parse_dt(lot.get("auction")))
    for k, title, width, kind in COLS:
        c = ws.cell(row, C[k])
        if k == "ppm":
            c.value = f'=IF(AND(ISNUMBER(L{row}),ISNUMBER(G{row}),G{row}>0),L{row}/G{row},"")'
        elif k == "days":
            c.value = f'=IF(ISNUMBER(P{row}),MAX(0,INT(P{row}-TODAY())),"")'
        elif kind == "link":
            if values.get("url"):
                c.value = "открыть ↗"
                c.hyperlink = values["url"]
        else:
            c.value = values.get(k)
            if c.value in (None, "") and (k in DASH or k == "step"):
                c.value = "—"
        c.font = font()
        if c.hyperlink:
            c.font = font(color=ACCENT)
        elif c.value == "—":
            c.font = font(color=GREY)
        c.border = Border(bottom=RULE, left=SEP if k in BLOCK_START else None)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True if kind in ("text", "ctext") else None)
        c.number_format = FORMATS.get(kind, "General")
    ws.cell(row, C["n"]).font = font(bold=True, color=ACCENT)
    ws.cell(row, C["address"]).font = font(bold=True)
    ws.cell(row, C["price"]).font = font(bold=True, size=11)
    if lot.get("price_src"):
        ws.cell(row, C["area"]).comment = Comment(lot["price_src"], "Источник")

table = Table(displayName="Reestr", ref=f"A1:{L(len(COLS))}{last}")
table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
ws.add_table(table)


def rng(key):
    return f"{L(C[key])}{first}:{L(C[key])}{last}"


ws.conditional_formatting.add(rng("okn"), CellIsRule(operator="equal", formula=['"Да"'], fill=fill("FFF1D6"),
                                                     font=Font(name=FONT, bold=True, color="9A5B00")))
ws.conditional_formatting.add(rng("ppm"), DataBarRule(start_type="num", start_value=0, end_type="max",
                                                      color=BAR, showValue=True))
q = f"{L(C['days'])}{first}"
for cond, bg, fg in ((f"{q}<=14", "FDE2E1", "B42318"),
                     (f"AND({q}>14,{q}<=30)", "FFF1D6", "9A5B00"),
                     (f"{q}>30", "E3F4EA", "1E7B45")):
    ws.conditional_formatting.add(rng("days"), FormulaRule(formula=[f"AND(ISNUMBER({q}),{cond})"],
                                                          fill=fill(bg), font=Font(name=FONT, bold=True, color=fg)))
# zebra as the lowest-priority rule, so stripes stay in place when the table is re-sorted
ws.conditional_formatting.add(f"A{first}:{L(len(COLS))}{last}",
                              FormulaRule(formula=["MOD(ROW(),2)=1"], fill=fill(ZEBRA)))

ws.freeze_panes = "D2"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "1:1"
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print(f"{OUT}: {len(lots)} lots")
