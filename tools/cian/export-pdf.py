#!/usr/bin/env python3
"""Собрать docs/premium-cao/premium-zhk-cao.pdf — вся книга одним документом под A3 (альбомная).

  python3 tools/cian/export-pdf.py

Читает готовый premium-zhk-cao.xlsx (значения, ширины колонок, ссылки), раскладывает
в HTML под ширину печатной полосы A3 и печатает Chromium'ом. Каждый раздел открывается
шапкой со сводкой (объектов, медиана, разброс), таблицы переносятся с повтором шапки,
карты идут во всю полосу. Шрифт — docs/premium-cao/fonts.css (tools/cian/fetch-fonts.py).
"""
import json, re, subprocess, html
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / 'docs' / 'premium-cao'
XLSX = DOCS / 'premium-zhk-cao.xlsx'
OUT = DOCS / 'premium-zhk-cao.pdf'
FONTS = DOCS / 'fonts.css'
HTML = Path('/tmp/claude-0/-home-user-test/5ef931c1-a5eb-584d-bea4-9b195abd4240/scratchpad/book.html')

# Основа — глубокий синий и графит; у каждого раздела свой цвет, и он же
# метит статус на картах и в списках. Всё остальное — оттенки серого.
INK, TEXT, MUTED, LINE, PAPER = '#0B1A33', '#1E2430', '#6B7280', '#DAD6CE', '#F7F4EE'
SECTION = {'1': ('#1F5C46', 'Построено'), '2': ('#B0642B', 'Строится'), '3': ('#3D5A80', 'Проектирование')}
STATUS = {'построено': '#1F5C46', 'строится': '#B0642B', 'проектирование': '#3D5A80'}

SHORT_HEAD = {'Корпусов': 'Корп.', 'Этажность': 'Этажей', 'Лотов в продаже': 'Лотов',
              'Площадь лотов, м²': 'Площадь, м²', 'Ссылка на Циан': 'Циан', 'Ссылка на проект': 'Проект',
              'Цена за метр ОТ': 'Цена ₽/м², от', 'Цена за метр медиана': 'Цена ₽/м², медиана',
              'Цена за метр ДО': 'Цена ₽/м², до', 'Общая площадь, м²': 'Площадь, м²',
              'Точный адрес / адреса': 'Точный адрес', 'Первое упоминание': 'Первое упом.',
              'Последнее упоминание': 'Последнее упом.', 'Достоверность': 'Достов.',
              'Срок сдачи': 'Сдача', 'Год постройки': 'Год', 'Год сдачи': 'Год', 'Цена от, ₽/м²': 'Цена ₽/м², от'}

def esc(v): return html.escape(str(v)) if v is not None else ''

def num(v): return f'{int(v):,}'.replace(',', ' ')   # обычный пробел: узкого в подмножестве шрифта нет

def fmt(v):
    if isinstance(v, bool) or v is None: return ''
    if isinstance(v, (int, float)): return num(v)
    return esc(v)

def median(xs): return sorted(xs)[len(xs) // 2] if xs else None

wb = load_workbook(XLSX)
maps_idx = json.load(open(DOCS / 'maps' / 'index.json')) if (DOCS / 'maps' / 'index.json').exists() else {'maps': []}
parts = []
period = None
for ws in wb:
    m = re.search(r'((?:январ|феврал|март|апрел|ма[йя]|июн|июл|август|сентябр|октябр|ноябр|декабр)\w*\s+\d{4})',
                  str(ws['A2'].value or ''))
    if m: period = m.group(1); break

def body_rows(ws):
    return [r for r in ws.iter_rows(min_row=4, max_row=ws.max_row) if r[0].value]

def is_group(row): return all(c.value is None for c in row[1:])

def metric(label, value):
    return f'<div class="m"><div class="m-l">{esc(label)}</div><div class="m-v">{value}</div></div>'

for ws in wb:
    title = str(ws['A1'].value or ws.title)
    subtitle = str(ws['A2'].value or '')
    sec = title.split('.')[0].strip() if re.match(r'^\d+\.', title) else ''
    color, kicker = SECTION.get(sec, (INK, 'Приложение'))
    name = re.sub(r'^\d+\.\s*', '', title)
    name = re.sub(r'^Карта\s*·\s*', '', name)   # слово «Карта» уже стоит над заголовком

    # ---------- карта ----------
    if ws.title.startswith('Карта'):
        mp = next((m for m in maps_idx['maps'] if m['slug'] in ws.title.lower()
                   or any(w in ws.title for w in m['title'].split(',')[0].split())), None)
        if not mp: continue
        legend = ''.join(f'<span class="lg"><i style="background:{c}"></i>{esc(s)}</span>' for s, c in STATUS.items())
        parts.append(f'''<section class="page map">
  <header class="ph"><div class="ph-l"><div class="kicker">Карта</div><h2>{esc(name)}</h2>
    <p class="sub">{esc(subtitle)}</p></div><div class="ph-r">{legend}</div></header>
  <div class="mapwrap"><img src="{(DOCS / mp['file']).as_uri()}"></div>
</section>''')
        continue

    # ---------- таблица ----------
    ncol = ws.max_column
    widths = [ws.column_dimensions[get_column_letter(i + 1)].width or 10 for i in range(ncol)]
    total = sum(widths)
    head = [str(c.value or '').strip() for c in ws[3]]
    body = body_rows(ws)
    data = [r for r in body if not is_group(r)]
    heat_i = next((i for i, h in enumerate(head) if 'медиана' in h or h.startswith('Цена от')), None)
    vals = [r[heat_i].value for r in data if heat_i is not None and isinstance(r[heat_i].value, (int, float))]
    lo, hi = (min(vals), max(vals)) if vals else (0, 1)

    # сводка раздела в шапке
    unit = ('ЖК' if sec in ('1', '2') else 'площадок' if sec == '3'
            else 'объектов' if ws.title.startswith('Список') else 'каналов' if ws.title.startswith('Источники') else 'строк')
    metrics = [metric('Всего', f'{len(data)} <small>{unit}</small>')]
    if vals:
        metrics.append(metric('Медиана, ₽/м²', num(median(vals))))
        metrics.append(metric('Разброс, ₽/м²', f'{num(lo)} – {num(hi)}'))
    if period: metrics.append(metric('Данные', esc(period)))

    cols = ''.join(f'<col style="width:{w / total * 100:.3f}%">' for w in widths)
    th = ''.join(f'<th>{esc(SHORT_HEAD.get(h, h))}</th>' for h in head)
    groups = [[None, []]]        # [заголовок кластера, строки]
    i = n = 0
    while i < len(body):
        row = body[i]
        if is_group(row):
            txt = str(row[0].value)
            m = re.match(r'^(.*?)\s*—\s*(\d+\s.*)$', txt)
            lbl, cnt = (m.group(1), m.group(2)) if m else (txt, '')
            # медиана кластера — по строкам до следующего заголовка
            j = i + 1; block = []
            while j < len(body) and not is_group(body[j]):
                v = body[j][heat_i].value if heat_i is not None else None
                if isinstance(v, (int, float)): block.append(v)
                j += 1
            extra = f'<span class="g-m">медиана <b>{num(median(block))}</b> ₽/м²</span>' if block else ''
            groups.append([f'<tr class="grp"><td colspan="{ncol}"><div class="gb" style="border-left-color:{color}">'
                           f'<span class="g-n">{esc(lbl)}</span><span class="g-c">{esc(cnt)}</span>{extra}</div></td></tr>', []])
            i += 1; n = 0
            continue
        n += 1
        tds = []
        for k, c in enumerate(row):
            v = c.value
            if c.hyperlink is not None or (isinstance(v, str) and v.startswith('http')):
                href = c.hyperlink.target if c.hyperlink is not None else v
                label = str(v).replace(' ↗', '').lower() if isinstance(v, str) and not v.startswith('http') else 'открыть'
                tds.append(f'<td class="lnk"><a href="{esc(href)}">{esc(label)}</a></td>'); continue
            if k == heat_i and isinstance(v, (int, float)):
                w = 6 + 94 * ((v - lo) / (hi - lo) if hi > lo else 0.5)
                tds.append(f'<td class="bar"><i style="width:{w:.1f}%;background:{color}"></i><b>{num(v)}</b></td>'); continue
            if isinstance(v, str) and v in STATUS:
                tds.append(f'<td class="st"><i style="background:{STATUS[v]}"></i>{esc(v)}</td>'); continue
            cls = 'nm' if k == 0 else ('num' if isinstance(v, (int, float)) else ('note' if k == ncol - 1 else ''))
            tds.append(f'<td class="{cls}">{fmt(v)}</td>')
        groups[-1][1].append(f'<tr class="{"odd" if n % 2 == 0 else ""}">' + ''.join(tds) + '</tr>')
        i += 1
    KEEP = 6           # кластер до 6 строк печатается целиком; у больших с заголовком держатся первые 2 строки
    tbodies = []
    for hdr, rows in groups:
        if not hdr and not rows: continue
        if not hdr: tbodies.append('<tbody>' + ''.join(rows) + '</tbody>'); continue
        if len(rows) <= KEEP: tbodies.append('<tbody class="keep">' + hdr + ''.join(rows) + '</tbody>')
        else: tbodies.append('<tbody class="keep">' + hdr + ''.join(rows[:2]) + '</tbody><tbody>' + ''.join(rows[2:]) + '</tbody>')

    parts.append(f'''<section class="page">
  <header class="ph" style="--c:{color}"><div class="ph-l"><div class="kicker">{esc(kicker)}</div>
    <h2>{esc(name)}</h2><p class="sub">{esc(subtitle)}</p></div>
    <div class="ph-r">{''.join(metrics)}</div></header>
  <table class="grid"><colgroup>{cols}</colgroup><thead><tr>{th}</tr></thead>{''.join(tbodies)}</table>
</section>''')

fonts = FONTS.read_text(encoding='utf-8') if FONTS.exists() else ''
css = fonts + f'''
* {{ box-sizing: border-box; }}
html, body {{ margin: 0; }}
body {{ font-family: 'IBM Plex Sans', 'Liberation Sans', Arial, sans-serif; color: {TEXT}; background: #fff;
  -webkit-font-smoothing: antialiased; font-variant-numeric: lining-nums tabular-nums; }}
section.page {{ page-break-after: always; }}
section.page:last-child {{ page-break-after: auto; }}

/* ---- шапка раздела: слева название, справа сводка ---- */
.ph {{ display: flex; justify-content: space-between; align-items: flex-end; gap: 12mm;
  border-bottom: .6px solid {INK}; padding-bottom: 8px; margin-bottom: 10px; }}
.ph-l {{ flex: 1; min-width: 0; }}
.kicker {{ font-size: 6.8pt; font-weight: 500; color: var(--c, {INK}); margin-bottom: 5px; letter-spacing: .06em; }}
.kicker::before {{ content: ''; display: inline-block; width: 5px; height: 5px; border-radius: 50%;
  background: var(--c, {INK}); margin-right: 6px; vertical-align: 1px; }}
h2 {{ font-family: Spectral, 'Liberation Serif', Georgia, serif; font-size: 24pt; font-weight: 500;
  margin: 0; letter-spacing: -.005em; line-height: 1; color: {INK}; }}
.sub {{ font-size: 7.4pt; color: {MUTED}; margin: 6px 0 0; line-height: 1.4; max-width: 190mm; }}
.ph-r {{ display: flex; gap: 10mm; flex-shrink: 0; align-items: flex-end; }}
.m-l {{ font-size: 6.2pt; color: {MUTED}; font-weight: 400; letter-spacing: .04em; }}
.m-v {{ font-family: Spectral, 'Liberation Serif', Georgia, serif; font-size: 14pt; font-weight: 500;
  color: {INK}; white-space: nowrap; margin-top: 2px; line-height: 1.05; letter-spacing: .01em; }}
.m-v small {{ font-family: 'IBM Plex Sans', sans-serif; font-size: 7pt; font-weight: 500; color: {MUTED}; }}

/* ---- таблица ---- */
table {{ width: 100%; border-collapse: collapse; table-layout: fixed; }}
table.grid {{ font-size: 6.6pt; }}
table.grid th {{ color: {MUTED}; font-size: 5.8pt; font-weight: 500; padding: 4px 3px 5px; background: #fff;
  border-bottom: .6px solid {INK}; letter-spacing: .02em; text-align: center; vertical-align: bottom; line-height: 1.15;
  word-wrap: break-word; overflow-wrap: anywhere; hyphens: auto; }}
table.grid td {{ border-bottom: .5px solid {LINE}; padding: 3px 3px; text-align: center; vertical-align: middle;
  line-height: 1.3; word-wrap: break-word; overflow-wrap: anywhere; hyphens: auto; }}
table.grid td.num {{ white-space: nowrap; }}
table.grid td.nm {{ font-weight: 500; font-size: 6.9pt; color: {INK}; }}
table.grid td.note {{ color: #3F4652; font-size: 6.1pt; line-height: 1.32; }}
table.grid tr.odd td {{ background: {PAPER}; }}
thead {{ display: table-header-group; }}
tr {{ page-break-inside: avoid; }}

td.bar {{ position: relative; padding: 0 3px; }}
td.bar i {{ position: absolute; left: 4%; top: 20%; height: 60%; opacity: .2; border-radius: 1px; }}
td.bar b {{ position: relative; font-weight: 500; color: {INK}; white-space: nowrap; }}
td.st {{ white-space: nowrap; }}
td.st i {{ display: inline-block; width: 5px; height: 5px; border-radius: 50%; margin-right: 4px; vertical-align: 1px; }}
td.lnk a {{ color: {INK}; font-size: 6pt; text-decoration: none; border-bottom: .5px solid #B5B0A6; }}

/* район — тёмная полоса во всю ширину, отделённая воздухом от предыдущего блока */
tr.grp td {{ background: #fff; border: 0; padding: 9px 0 0; text-align: left; }}
tbody.keep + tbody tr:first-child td, tbody + tbody.keep tr.grp td {{ }}
.gb {{ background: {INK}; color: #fff; padding: 3.5px 10px 3.5px 11px; border-left: 3px solid; display: flex;
  align-items: baseline; gap: 10px; }}
.g-n {{ font-family: Spectral, 'Liberation Serif', Georgia, serif; font-size: 11pt; font-weight: 500;
  letter-spacing: .02em; }}
.g-c {{ font-size: 6.6pt; color: #C9D1E0; }}
.g-m {{ font-size: 6.6pt; color: #C9D1E0; margin-left: auto; }}
.g-m b {{ color: #fff; font-weight: 600; }}
tbody.keep {{ break-inside: avoid; page-break-inside: avoid; }}
section.compact table.grid td {{ padding: 1.6px 3px; }}
section.compact tr.grp td {{ padding-top: 4px; }}
section.compact .gb {{ padding: 1.5px 9px 1.5px 10px; }}
section.compact .g-n {{ font-size: 10pt; }}
section.compact .ph {{ margin-bottom: 7px; }}

/* ---- карта ---- */
section.map img {{ display: block; margin: 0 auto; break-inside: avoid; }}
.lg {{ font-size: 7pt; color: {TEXT}; margin-left: 9px; }}
.lg i {{ display: inline-block; width: 8px; height: 8px; border-radius: 50%; margin-right: 4px; vertical-align: -1px; }}
.mapwrap {{ margin-top: 3px; }}
.mapwrap img {{ max-width: 100%; max-height: 240mm; object-fit: contain; border: .5px solid {LINE}; }}
'''
HTML.write_text(f'<!doctype html><html lang="ru"><head><meta charset="utf-8"><style>{css}</style></head>'
                f'<body>{"".join(parts)}</body></html>', encoding='utf-8')

foot = (f'<div style="width:100%;font-family:\'IBM Plex Sans\',Arial,sans-serif;font-size:7px;color:#8A8F98;'
        f'padding:0 9mm;display:flex;justify-content:space-between;">'
        f'<span>Премиум-ЖК центра Москвы{" · " + period if period else ""}</span>'
        f'<span><span class="pageNumber"></span> / <span class="totalPages"></span></span></div>')
js = f'''
const {{ chromium }} = require('/opt/node22/lib/node_modules/playwright');
(async () => {{
  const b = await chromium.launch({{ executablePath: '/opt/pw-browsers/chromium-1194/chrome-linux/chrome', headless: true, args: ['--no-sandbox'] }});
  const p = await (await b.newContext()).newPage();
  p.on('console', (m) => console.log('  ' + m.text()));
  await p.goto('file://{HTML}', {{ waitUntil: 'networkidle', timeout: 240000 }});
  await p.evaluate(() => document.fonts.ready);
  await p.emulateMedia({{ media: 'print' }});
  await p.setViewportSize({{ width: 1528, height: 1051 }});   // печатная полоса A3: 404 × 278 мм
  await p.evaluate(() => {{
    const PAGE = 278 * 96 / 25.4;
    for (const s of document.querySelectorAll('section.page')) {{
      const head = s.querySelector('.ph');
      const used = head.getBoundingClientRect().height + parseFloat(getComputedStyle(head).marginBottom);
      if (s.classList.contains('map')) {{
        s.querySelector('img').style.maxHeight = Math.floor(PAGE - used - 14) + 'px';
        continue;
      }}
      /* Порядок попыток: как есть → плотные строки → плотные строки и лёгкое ужатие (до 12%).
         Раздел, который всё равно не влезает в лист, печатается плотно, только если это
         убирает страницу: повтор шапки и неделимые строки съедают часть каждой следующей. */
      const H = () => s.getBoundingClientRect().height;
      const pages = () => H() <= PAGE - 8 ? 1 : 1 + Math.ceil((H() - (PAGE - 8)) / (PAGE - 40));
      if (H() > PAGE - 8) {{
        const before = pages();
        s.classList.add('compact');
        if (H() <= PAGE - 8) {{}}                                    // влезло плотными строками
        else if (H() <= PAGE * 1.12) s.style.zoom = ((PAGE - 3) / H()).toFixed(4);
        else if (pages() >= before) s.classList.remove('compact');   // многостраничный: плотно не помогло
      }}
      const h = H();
      console.log(`${{s.querySelector('h2').textContent}}: ${{(h / PAGE).toFixed(2)}} стр.` +
                  (s.classList.contains('compact') ? ' плотно' : '') + (s.style.zoom ? ' zoom ' + s.style.zoom : ''));
    }}
  }});
  await p.pdf({{ path: '{OUT}', format: 'A3', landscape: true, printBackground: true,
    displayHeaderFooter: true, headerTemplate: '<div></div>', footerTemplate: {json.dumps(foot)},
    margin: {{ top: '8mm', bottom: '11mm', left: '8mm', right: '8mm' }} }});
  await b.close();
}})();
'''
jsf = HTML.with_suffix('.js'); jsf.write_text(js)
subprocess.run(['node', str(jsf)], check=True)
print(f'-> {OUT} ({OUT.stat().st_size // 1024} КБ)')
